"""Batch data preparation for local model training.

Usage:
    python -m scripts.prepare_training_data \
        --input-dir data/paired_docs \
        --output-dir training_data

Input directory structure (expected):
    data/paired_docs/
    ├── 项目A/
    │   ├── tender.docx     # 招标文件
    │   └── bid.docx        # 投标文件
    ├── 项目B/
    │   ├── 招标文件.docx
    │   └── 投标文件.docx
    ...

Outputs:
    training_data/
    ├── structure_pairs.jsonl     # 结构提取训练数据 (input: 招标原文, output: JSON)
    ├── narrative_chunks.jsonl    # 叙述段落语料 (title + hints → content)
    ├── rag_corpus.jsonl          # RAG 向量库语料 (所有段落)
    └── stats.json                # 处理统计信息
"""

import argparse
import json
import os
import re
import sys
from pathlib import Path
from typing import Dict, List, Optional, Tuple

# Add project root to path
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from docx import Document


def parse_docx(path: str) -> Dict:
    """Parse a .docx file and extract sections (standalone version of tender_parsing)."""
    doc = Document(path)
    sections = []
    current_order = 0
    current_title = ""
    current_level = 0
    current_content_parts = []
    current_tables = []
    full_text_parts = []

    def _flush_section():
        nonlocal current_order, current_title, current_level
        nonlocal current_content_parts, current_tables
        if current_title:
            content = "\n".join(current_content_parts).strip()
            has_tables = len(current_tables) > 0
            sec_type = _guess_type(current_title, content, has_tables)
            sections.append({
                "order": current_order,
                "title": current_title,
                "level": current_level,
                "content": content,
                "section_type": sec_type,
                "tables": current_tables,
            })
            current_content_parts = []
            current_tables = []

    def _extract_table(table) -> Dict:
        rows = []
        headers = []
        for i, row in enumerate(table.rows):
            cells = [cell.text.strip() for cell in row.cells]
            if i == 0:
                headers = cells
            rows.append(cells)
        return {"headers": headers, "rows": rows}

    cn_heading_patterns = [
        r'^[一二三四五六七八九十]+[、.]',
        r'^（[一二三四五六七八九十]+）',
        r'^第[一二三四五六七八九十]+[章节部分]',
        r'^\d+[、.\s]',
    ]

    for element in doc.element.body:
        if element.tag.endswith('}p'):
            for para in doc.paragraphs:
                if para._element is element:
                    text = para.text.strip()
                    full_text_parts.append(text)
                    style_name = para.style.name if para.style else ""
                    is_heading = style_name.startswith("Heading")

                    if not is_heading and text:
                        for pattern in cn_heading_patterns:
                            if re.match(pattern, text):
                                is_heading = True
                                break

                    if is_heading and text:
                        _flush_section()
                        if style_name.startswith("Heading"):
                            try:
                                current_level = int(style_name.split()[-1])
                            except (ValueError, IndexError):
                                current_level = 1
                        else:
                            current_level = 2
                        current_order += 1
                        current_title = text
                    elif text:
                        current_content_parts.append(text)
                    break

        elif element.tag.endswith('}tbl'):
            for table in doc.tables:
                if table._element is element:
                    table_data = _extract_table(table)
                    current_tables.append(table_data)
                    for row in table_data["rows"]:
                        full_text_parts.append(" | ".join(row))
                    break

    _flush_section()

    if not sections:
        full_text = "\n".join(full_text_parts)
        sections.append({
            "order": 1, "title": "正文", "level": 1,
            "content": full_text, "section_type": "narrative", "tables": []
        })

    return {
        "raw_text": "\n".join(full_text_parts),
        "sections": sections,
        "total_sections": len(sections),
    }


def _guess_type(title: str, content: str, has_tables: bool) -> str:
    """Heuristic section type classification."""
    table_kw = ["一览表", "清单", "明细表", "情况表", "统计表", "报价表"]
    if any(kw in title for kw in table_kw) or has_tables:
        return "table"
    qual_kw = ["资质", "证书", "执业", "许可", "授权", "委托", "声明函", "承诺函",
               "营业执照", "审计报告"]
    if any(kw in title for kw in qual_kw):
        return "qualification"
    form_kw = ["投标函", "法定代表人"]
    if any(kw in title for kw in form_kw):
        return "form"
    return "narrative"


def find_doc_pairs(input_dir: str) -> List[Tuple[str, str, str]]:
    """Find paired tender/bid documents in subdirectories.

    Returns: List of (project_name, tender_path, bid_path)
    """
    pairs = []
    tender_patterns = ["tender", "招标", "zbwj", "招标文件"]
    bid_patterns = ["bid", "投标", "tbwj", "投标文件", "应答"]

    input_path = Path(input_dir)

    # Case 1: Subdirectories per project
    for subdir in sorted(input_path.iterdir()):
        if not subdir.is_dir():
            continue

        docx_files = list(subdir.glob("*.docx"))
        if len(docx_files) < 2:
            continue

        tender_file = None
        bid_file = None

        for f in docx_files:
            name_lower = f.stem.lower()
            if any(p in name_lower for p in tender_patterns):
                tender_file = str(f)
            elif any(p in name_lower for p in bid_patterns):
                bid_file = str(f)

        # If only 2 files and can't tell, assume first=tender second=bid
        if not tender_file and not bid_file and len(docx_files) == 2:
            tender_file = str(docx_files[0])
            bid_file = str(docx_files[1])

        if tender_file and bid_file:
            pairs.append((subdir.name, tender_file, bid_file))

    # Case 2: Flat directory with naming convention
    if not pairs:
        docx_files = sorted(input_path.glob("*.docx"))
        # Try to find pairs by matching names
        tender_files = {}
        bid_files = {}
        for f in docx_files:
            name = f.stem.lower()
            if any(p in name for p in tender_patterns):
                # Extract project identifier
                key = re.sub(r'(tender|招标|zbwj|招标文件)', '', name).strip('_- ')
                tender_files[key] = str(f)
            elif any(p in name for p in bid_patterns):
                key = re.sub(r'(bid|投标|tbwj|投标文件|应答)', '', name).strip('_- ')
                bid_files[key] = str(f)

        for key in tender_files:
            if key in bid_files:
                pairs.append((key or f"project_{len(pairs)+1}",
                               tender_files[key], bid_files[key]))

    return pairs


def generate_structure_pair(tender_result: Dict, bid_result: Dict,
                            project_name: str) -> Dict:
    """Generate one structure extraction training sample.

    Input: tender raw text
    Output: structured JSON derived from the ACTUAL bid document
    """
    # Build output JSON from bid document's actual structure
    volumes = []
    current_volume = {"name": "投标文件", "sections": []}

    for sec in bid_result["sections"]:
        # Detect volume boundaries (level 1 headings with 分册 keyword)
        if sec["level"] == 1 and ("分册" in sec["title"] or "部分" in sec["title"]):
            if current_volume["sections"]:
                volumes.append(current_volume)
            current_volume = {"name": sec["title"], "sections": []}
            continue

        current_volume["sections"].append({
            "order": sec["order"],
            "title": sec["title"],
            "type": sec["section_type"],
            "required": True,
            "content_hints": sec["content"][:100] if sec["content"] else "",
            "data_fields": [],
        })

    if current_volume["sections"]:
        volumes.append(current_volume)

    output_json = {
        "bid_title": f"{project_name} 投标文件",
        "volumes": volumes,
        "qualification_requirements": [],
        "format_requirements": {},
        "evaluation_criteria": [],
        "deadline_info": {},
    }

    return {
        "instruction": "请分析以下招标文件内容，提取投标文件需要包含的所有章节和要求。输出 JSON 格式。",
        "input": tender_result["raw_text"][:12000],  # Truncate for model context
        "output": json.dumps(output_json, ensure_ascii=False, indent=2),
        "project": project_name,
    }


def extract_narrative_chunks(bid_result: Dict, project_name: str) -> List[Dict]:
    """Extract narrative sections from bid documents for RAG/training."""
    chunks = []
    for sec in bid_result["sections"]:
        if sec["section_type"] != "narrative":
            continue
        content = sec.get("content", "").strip()
        if len(content) < 50:
            continue

        chunks.append({
            "title": sec["title"],
            "content": content,
            "project": project_name,
            "type": "narrative",
            "length": len(content),
        })

    return chunks


def extract_rag_corpus(bid_result: Dict, project_name: str) -> List[Dict]:
    """Extract all chunks for RAG knowledge base."""
    chunks = []
    for sec in bid_result["sections"]:
        content = sec.get("content", "").strip()
        if len(content) < 30:
            continue

        chunks.append({
            "title": sec["title"],
            "content": content,
            "project": project_name,
            "section_type": sec["section_type"],
        })

        # Also include table content as separate chunks
        for table in sec.get("tables", []):
            if table.get("rows"):
                table_text = "\n".join(" | ".join(row) for row in table["rows"])
                chunks.append({
                    "title": f"{sec['title']} (表格)",
                    "content": table_text,
                    "project": project_name,
                    "section_type": "table_data",
                })

    return chunks


def main():
    parser = argparse.ArgumentParser(description="Prepare training data from paired tender/bid documents")
    parser.add_argument("--input-dir", required=True, help="Directory containing paired documents")
    parser.add_argument("--output-dir", default="training_data", help="Output directory (default: training_data)")
    args = parser.parse_args()

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    # Find document pairs
    print(f"🔍 Scanning {args.input_dir} for paired documents...")
    pairs = find_doc_pairs(args.input_dir)
    print(f"   Found {len(pairs)} pairs\n")

    if not pairs:
        print("❌ No paired documents found.")
        print("Expected structure:")
        print("  input_dir/项目A/招标文件.docx + 投标文件.docx")
        print("  input_dir/项目B/tender.docx + bid.docx")
        sys.exit(1)

    # Process each pair
    structure_pairs = []
    narrative_chunks = []
    rag_corpus = []
    stats = {"total_pairs": len(pairs), "processed": 0, "errors": [], "section_types": {}}

    for i, (name, tender_path, bid_path) in enumerate(pairs, 1):
        print(f"[{i}/{len(pairs)}] Processing: {name}")
        print(f"   Tender: {Path(tender_path).name}")
        print(f"   Bid:    {Path(bid_path).name}")

        try:
            # Parse both documents
            tender_result = parse_docx(tender_path)
            bid_result = parse_docx(bid_path)

            print(f"   Tender: {tender_result['total_sections']} sections, "
                  f"{len(tender_result['raw_text'])} chars")
            print(f"   Bid:    {bid_result['total_sections']} sections, "
                  f"{len(bid_result['raw_text'])} chars")

            # Generate structure extraction training pair
            pair = generate_structure_pair(tender_result, bid_result, name)
            structure_pairs.append(pair)

            # Extract narrative chunks
            narr = extract_narrative_chunks(bid_result, name)
            narrative_chunks.extend(narr)
            print(f"   Narrative chunks: {len(narr)}")

            # Extract RAG corpus
            rag = extract_rag_corpus(bid_result, name)
            rag_corpus.extend(rag)

            # Count section types
            for sec in bid_result["sections"]:
                t = sec["section_type"]
                stats["section_types"][t] = stats["section_types"].get(t, 0) + 1

            stats["processed"] += 1

        except Exception as e:
            print(f"   ❌ Error: {e}")
            stats["errors"].append({"project": name, "error": str(e)})

        print()

    # Write outputs
    structure_path = output_dir / "structure_pairs.jsonl"
    with open(structure_path, "w", encoding="utf-8") as f:
        for item in structure_pairs:
            f.write(json.dumps(item, ensure_ascii=False) + "\n")

    narrative_path = output_dir / "narrative_chunks.jsonl"
    with open(narrative_path, "w", encoding="utf-8") as f:
        for item in narrative_chunks:
            f.write(json.dumps(item, ensure_ascii=False) + "\n")

    rag_path = output_dir / "rag_corpus.jsonl"
    with open(rag_path, "w", encoding="utf-8") as f:
        for item in rag_corpus:
            f.write(json.dumps(item, ensure_ascii=False) + "\n")

    stats_path = output_dir / "stats.json"
    stats.update({
        "structure_pairs": len(structure_pairs),
        "narrative_chunks": len(narrative_chunks),
        "rag_chunks": len(rag_corpus),
    })
    with open(stats_path, "w", encoding="utf-8") as f:
        json.dump(stats, f, ensure_ascii=False, indent=2)

    # Summary
    print("=" * 60)
    print("📊 处理完成")
    print(f"   成功: {stats['processed']}/{stats['total_pairs']}")
    print(f"   结构提取训练数据: {len(structure_pairs)} 对 → {structure_path}")
    print(f"   叙述段落:        {len(narrative_chunks)} 条 → {narrative_path}")
    print(f"   RAG 语料:        {len(rag_corpus)} 条 → {rag_path}")
    print(f"   统计:            {stats_path}")
    if stats["errors"]:
        print(f"   ⚠️ 错误: {len(stats['errors'])} 个")
    print()
    print("下一步:")
    print(f"  1. 检查 {structure_path} 中的 JSON 结构是否准确")
    print(f"  2. 人工校对后运行 train_structure_model.py")


if __name__ == "__main__":
    main()
