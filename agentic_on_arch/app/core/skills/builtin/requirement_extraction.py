"""Requirement extraction skill — analyze tender → derive bid document structure.

V3: Multi-pass LLM analysis.
  Pass 1: Understand the tender, extract key requirements & rejection conditions
  Pass 2: Generate bid document structure from the analysis
"""

import json
from typing import Any, Dict, List

from app.core.skills.base import BaseSkill
from app.core.llm import get_llm
from app.utils.logger import logger

# ── Prompts are now managed separately in app/core/prompts/ ──
# To modify prompt behavior, edit: app/core/prompts/requirement_prompts.py
from app.core.prompts.requirement_prompts import (
    ANALYSIS_SYSTEM,
    ANALYSIS_PROMPT,
    STRUCTURE_SYSTEM,
    STRUCTURE_PROMPT,
    BATCH_CLASSIFY_SYSTEM,
    BATCH_CLASSIFY_PROMPT,
)




def _safe_parse_json(text):
    # type: (str) -> Any
    """Try to parse JSON from LLM output, handling common issues."""
    text = text.strip()

    # Remove markdown code fences if present
    if text.startswith("```"):
        lines = text.split("\n")
        if lines[0].startswith("```"):
            lines = lines[1:]
        if lines and lines[-1].strip() == "```":
            lines = lines[:-1]
        text = "\n".join(lines)

    try:
        return json.loads(text)
    except json.JSONDecodeError:
        import re
        for pattern in [r'\{[\s\S]*\}', r'\[[\s\S]*\]']:
            match = re.search(pattern, text)
            if match:
                try:
                    return json.loads(match.group())
                except json.JSONDecodeError:
                    continue
        logger.error(f"Failed to parse JSON from LLM output: {text[:200]}...")
        return None


class RequirementExtractionSkill(BaseSkill):
    """Analyze tender document and derive bid document structure.

    V3 approach (default):
      Pass 1: Deep analysis of tender — extract requirements, rejection
              conditions, evaluation criteria, composition rules
      Pass 2: Generate bid document structure from the analysis

    Falls back to V2 (heading classification) if V3 fails.
    """

    name = "requirement_extraction"
    description = "分析招标文件，提取废标项和评分标准，智能生成投标文件目录结构"

    MAX_CHUNK_SIZE = 28000  # Qwen-Max supports ~32K tokens

    async def execute(self, params):
        # type: (Dict[str, Any]) -> Any
        """
        Params:
            raw_text (str): Full text of the tender document
            sections (List[Dict]): Parsed sections from tender_parsing
            llm_provider (str): Which LLM to use (default: "qwen")
            mode (str): 'analyze' (V3, default) or 'structure' (V2) or 'legacy' (V1)
        """
        raw_text = params.get("raw_text", "")
        sections = params.get("sections", [])
        llm_provider = params.get("llm_provider", "qwen")
        mode = params.get("mode", "analyze")

        if not raw_text and not sections:
            raise ValueError("Either raw_text or sections must be provided")

        llm = get_llm(llm_provider)
        logger.info(f"Extracting requirements using {llm.get_model_name()}, mode={mode}")

        if mode == "analyze":
            # ── V3: Multi-pass tender analysis ──
            progress_cb = params.get("progress_callback", None)
            tender_filename = params.get("tender_filename", "")
            result = await self._analyze_and_build(
                llm, raw_text, sections,
                progress_cb=progress_cb,
                tender_filename=tender_filename,
            )
        elif mode == "structure" and sections:
            # ── V2: Classify parsed section titles ──
            result = await self._structure_from_sections(llm, sections, raw_text)
        else:
            # ── V1 Legacy ──
            result = await self._extract_full(llm, raw_text)

        if result is None:
            result = self._fallback_from_sections(sections)

        # Post-LLM: refine section types using local classifier
        result = self._refine_section_types(result)

        total_sections = sum(len(v.get('sections', [])) for v in result.get('volumes', []))
        logger.info(f"Extraction complete: {len(result.get('volumes', []))} volumes, "
                     f"{total_sections} sections (mode={mode})")
        return result

    # ── V3: Multi-pass analysis ──

    # ── Lot/package detection patterns ──
    _LOT_PATTERNS = [
        # “一标段” “第一标段” “1标段”
        ('\u4e00标段', '标段一'), ('\u7b2c\u4e00\u6807\u6bb5', '标段一'),
        ('\u4e8c标段', '标段二'), ('\u7b2c\u4e8c\u6807\u6bb5', '标段二'),
        ('\u4e09标段', '标段三'), ('\u7b2c\u4e09\u6807\u6bb5', '标段三'),
        ('1标段', '标段一'), ('2标段', '标段二'), ('3标段', '标段三'),
    ]

    @staticmethod
    def _detect_lot_info(filename, raw_text_head=''):
        # type: (str, str) -> Dict[str, Any]
        """Detect lot/package info from filename and first ~500 chars of text.
        Returns dict with lot_name (e.g. '标段一') or empty dict."""
        search_text = f"{filename} {raw_text_head[:500]}"
        for pattern, lot_name in RequirementExtractionSkill._LOT_PATTERNS:
            if pattern in search_text:
                return {"lot_name": lot_name, "source": "filename" if pattern in filename else "text"}
        return {}

    SCORING_KEYWORDS = [
        '评标', '评审', '评分', '打分', '计分',
        '评标办法', '评标标准', '评分标准', '评分细则',
        '评审标准', '评审办法', '评分方法',
    ]

    REJECTION_KEYWORDS = [
        '废标', '否决', '无效投标', '不予受理', '拒绝',
        '投标人须知', '须知前附表', '资格条件', '资格要求',
    ]

    @staticmethod
    def _render_table_as_markdown(table_data):
        # type: (Dict) -> str
        """Render a parsed table (from tender_parsing) as a markdown table."""
        rows = table_data.get('rows', [])
        if not rows:
            return ''
        # Build markdown table
        lines = []
        for i, row in enumerate(rows):
            line = '| ' + ' | '.join(cell.replace('|', '/') for cell in row) + ' |'
            lines.append(line)
            if i == 0:
                # Add separator after header
                lines.append('|' + '|'.join(['---'] * len(row)) + '|')
        return '\n'.join(lines)

    def _extract_scoring_sections(self, sections):
        # type: (List[Dict]) -> str
        """Locate and extract scoring/evaluation sections with table data."""
        scoring_parts = []
        for sec in (sections or []):
            title = sec.get('title', '')
            if any(kw in title for kw in self.SCORING_KEYWORDS):
                part = f"## {title}"
                content = sec.get('content', '')
                if content and len(content) > 20:
                    part += f"\n{content}"
                # Extract tables within this section
                for tbl in sec.get('tables', []):
                    md_table = self._render_table_as_markdown(tbl)
                    if md_table:
                        part += f"\n\n{md_table}"
                scoring_parts.append(part)
        return '\n\n'.join(scoring_parts) if scoring_parts else ''

    def _extract_rejection_sections(self, sections):
        # type: (List[Dict]) -> str
        """Locate and extract rejection/disqualification sections with table data."""
        rejection_parts = []
        for sec in (sections or []):
            title = sec.get('title', '')
            if any(kw in title for kw in self.REJECTION_KEYWORDS):
                part = f"## {title}"
                content = sec.get('content', '')
                if content and len(content) > 20:
                    part += f"\n{content}"
                for tbl in sec.get('tables', []):
                    md_table = self._render_table_as_markdown(tbl)
                    if md_table:
                        part += f"\n\n{md_table}"
                rejection_parts.append(part)
        return '\n\n'.join(rejection_parts) if rejection_parts else ''

    async def _analyze_and_build(self, llm, raw_text, sections, progress_cb=None, tender_filename=''):
        # type: (Any, str, List[Dict], Any, str) -> Dict
        """V3: Analyze tender → derive bid structure in two passes."""

        def _progress(msg):
            """Send progress update if callback is provided."""
            if progress_cb:
                try:
                    progress_cb(msg)
                except Exception:
                    pass

        async def _stream_with_progress(prompt, system, pass_name=""):
            """Use LLM streaming to show real-time generation progress."""
            buffer = []
            total_chars = 0
            last_report = 0
            REPORT_INTERVAL = 300  # Report every 300 chars

            try:
                async for chunk in llm.stream(prompt, system=system):
                    buffer.append(chunk)
                    total_chars += len(chunk)
                    # Send periodic progress with content snippet
                    if total_chars - last_report >= REPORT_INTERVAL:
                        last_report = total_chars
                        # Show last ~60 chars as a preview
                        recent = ''.join(buffer)[-80:].replace('\n', ' ').strip()
                        if len(recent) > 60:
                            recent = '...' + recent[-60:]
                        _progress(f"   └─ [{pass_name}] 已接收 {total_chars} 字符 | {recent}")
            except Exception as e:
                _progress(f"⚠️ [{pass_name}] 流式调用异常: {str(e)[:80]}")
                raise

            full_text = ''.join(buffer)
            _progress(f"   └─ [{pass_name}] 完成接收: 共 {len(full_text)} 字符")
            return full_text

        # ── Step 0: Extract scoring & rejection sections (with tables) ──
        scoring_text = self._extract_scoring_sections(sections)
        rejection_text = self._extract_rejection_sections(sections)
        if scoring_text:
            logger.info(f"Located scoring sections: {len(scoring_text)} chars "
                        f"(includes structured tables)")
            _progress(f"🏆 提取评分章节: {len(scoring_text)} 字符 (含表格结构)")
        if rejection_text:
            logger.info(f"Located rejection sections: {len(rejection_text)} chars")
            _progress(f"🔴 提取废标章节: {len(rejection_text)} 字符")

        # Prepare text: use full raw_text, truncate if too long
        text_for_analysis = raw_text
        if len(text_for_analysis) > self.MAX_CHUNK_SIZE:
            # Prioritize: keep beginning (project info, 投标人须知)
            # and end (投标文件格式, 评标办法 are often at the end)
            half = self.MAX_CHUNK_SIZE // 2
            text_for_analysis = (
                raw_text[:half] +
                "\n\n... (中间部分省略) ...\n\n" +
                raw_text[-half:]
            )
            logger.info(f"Tender text truncated: {len(raw_text)} → {len(text_for_analysis)} chars")

        # Append scoring sections (with structured tables) if truncated
        if scoring_text and len(raw_text) > self.MAX_CHUNK_SIZE:
            if scoring_text[:100] not in text_for_analysis:
                supplement = f"\n\n【评标办法原文（定向提取，含表格结构）】\n{scoring_text[:6000]}"
                text_for_analysis += supplement
                logger.info(f"Appended scoring section: +{len(supplement)} chars")

        # Append rejection sections if truncated
        if rejection_text and len(raw_text) > self.MAX_CHUNK_SIZE:
            if rejection_text[:100] not in text_for_analysis:
                supplement = f"\n\n【废标/否决条件原文（定向提取）】\n{rejection_text[:4000]}"
                text_for_analysis += supplement
                logger.info(f"Appended rejection section: +{len(supplement)} chars")

        # ── Step 0b: Detect lot info from filename ──
        lot_info = self._detect_lot_info(tender_filename, raw_text[:500])
        if lot_info:
            lot_name = lot_info['lot_name']
            logger.info(f"Lot detected: {lot_name} (source: {lot_info['source']})")
            _progress(f"🎯 检测到标段信息: {lot_name}")

        # ── Pass 1: Deep analysis (with retry) ──
        logger.info("Pass 1: Analyzing tender document...")
        _progress("🧠 Pass 1/3: AI 正在分析招标文件要求（可能需要 2-5 分钟）...")
        analysis = None
        try:
            prompt1 = ANALYSIS_PROMPT.format(tender_text=text_for_analysis)
            # Inject lot-specific hint if detected
            if lot_info:
                lot_hint = (
                    f"\n\n【重要提示：本次投标为{lot_info['lot_name']}】\n"
                    f"本招标文件可能包含多个标段的评分表。请只提取{lot_info['lot_name']}的评分标准。\n"
                    f"不同标段的评分标准可能不同（如标段一\"分所覆盖\"=国内分所，标段二=境外办公室）。\n"
                    f"evaluation_criteria 中只包含{lot_info['lot_name']}的评分项。"
                )
                prompt1 += lot_hint
                logger.info(f"Injected lot hint for {lot_info['lot_name']}")
            _progress(f"   └─ Prompt 大小: {len(prompt1)} 字符, 模型: {llm.get_model_name()}")
            # Try up to 2 times to handle transient Qwen API timeouts
            for attempt in range(2):
                try:
                    response1 = await _stream_with_progress(prompt1, system=ANALYSIS_SYSTEM, pass_name="Pass1")
                    analysis = _safe_parse_json(response1)
                    if analysis:
                        break
                except Exception as retry_err:
                    if attempt == 0:
                        logger.warning(f"Pass 1 attempt 1 failed ({retry_err}), retrying in 3s...")
                        _progress(f"⚠️ Pass 1 第1次尝试失败: {str(retry_err)[:80]}，3秒后重试...")
                        import asyncio
                        await asyncio.sleep(3)
                    else:
                        raise
            if analysis:
                # Log key findings
                rej_count = len(analysis.get("rejection_conditions", []))
                eval_count = len(analysis.get("evaluation_criteria", []))
                doc_count = len(analysis.get("bid_composition", {}).get("required_documents", []))
                has_format = analysis.get("bid_composition", {}).get("has_explicit_format", False)
                logger.info(
                    f"Pass 1 results: {doc_count} required docs, "
                    f"{rej_count} rejection conditions, "
                    f"{eval_count} evaluation criteria, "
                    f"explicit_format={'yes' if has_format else 'no'}"
                )
                _progress(
                    f"✅ Pass 1 完成: {doc_count} 个必须文件, "
                    f"{rej_count} 个废标条件, "
                    f"{eval_count} 个评分大类"
                )
        except Exception as e:
            logger.error(f"Pass 1 analysis failed: {e}")
            _progress(f"❌ Pass 1 失败: {str(e)[:100]}")

        if not analysis:
            logger.warning("Pass 1 failed, falling back to V2 section classification")
            _progress("⚠️ AI 分析失败，降级到规则分类模式...")
            if sections:
                return await self._structure_from_sections(llm, sections, raw_text)
            return None

        # ── Pass 2: Generate bid structure ──
        logger.info("Pass 2: Generating bid document structure...")
        _progress("📝 Pass 2/3: AI 正在生成投标文件大纲（可能需要 3-5 分钟）...")
        try:
            analysis_json = json.dumps(analysis, ensure_ascii=False, indent=2)
            # Provide extra tender context for Pass 2
            tender_context = raw_text[:6000] if raw_text else "暂无原文"

            prompt2 = STRUCTURE_PROMPT.format(
                analysis_json=analysis_json,
                tender_context=tender_context,
            )
            response2 = await _stream_with_progress(prompt2, system=STRUCTURE_SYSTEM, pass_name="Pass2")
            structure = _safe_parse_json(response2)

            if structure and structure.get("volumes"):
                # Attach analysis results to the structure for frontend use
                structure["tender_analysis"] = {
                    "project_info": analysis.get("project_info", {}),
                    "rejection_conditions": analysis.get("rejection_conditions", []),
                    "evaluation_criteria": analysis.get("evaluation_criteria", []),
                    "qualification_requirements": analysis.get("qualification_requirements", []),
                    "format_requirements": analysis.get("format_requirements", {}),
                    "deadline_info": analysis.get("deadline_info", {}),
                }
                # Ensure all required fields exist
                if "rejection_items" not in structure:
                    structure["rejection_items"] = []
                if "qualification_requirements" not in structure:
                    structure["qualification_requirements"] = analysis.get("qualification_requirements", [])
                if "format_requirements" not in structure:
                    structure["format_requirements"] = analysis.get("format_requirements", {})
                if "evaluation_criteria" not in structure:
                    structure["evaluation_criteria"] = analysis.get("evaluation_criteria", [])
                if "deadline_info" not in structure:
                    structure["deadline_info"] = analysis.get("deadline_info", {})

                # Count rejection risk sections
                total_secs = 0
                rej_secs = 0
                for vol in structure.get("volumes", []):
                    for sec in vol.get("sections", []):
                        total_secs += 1
                        if sec.get("rejection_risk"):
                            rej_secs += 1
                logger.info(f"Pass 2 results: {total_secs} sections, "
                             f"{rej_secs} with rejection risk")
                _progress(f"✅ Pass 2 完成: {total_secs} 个章节, {rej_secs} 个有废标风险")

                # ── Pass 2b: Sanitize catch-all sections ──
                # LLM sometimes generates vague sections like "商务部分" or
                # "综合实力" that overlap with specific sections. Remove them.
                _BANNED_PATTERNS = [
                    "商务部分", "综合实力", "律所综合实力",
                    "技术部分", "其他文件和资料", "补充材料",
                    "附件", "其他资料",
                ]
                for vol in structure.get("volumes", []):
                    original_count = len(vol.get("sections", []))
                    vol["sections"] = [
                        sec for sec in vol.get("sections", [])
                        if not any(
                            ban in sec.get("title", "")
                            for ban in _BANNED_PATTERNS
                        )
                    ]
                    removed = original_count - len(vol["sections"])
                    if removed > 0:
                        logger.info(
                            f"Pass 2b: Removed {removed} catch-all sections "
                            f"(banned patterns: {_BANNED_PATTERNS[:3]}...)"
                        )
                        _progress(f"🧹 移除 {removed} 个模糊章节")
                        # Re-number remaining sections
                        for i, sec in enumerate(vol["sections"], 1):
                            sec["order"] = i

                # ── Pass 3: Deterministic verification ──
                _progress("🔍 Pass 3/3: 校验 → 补全 → 联动匹配...")
                verification = self._verify_structure(analysis, structure)
                structure["verification"] = verification
                logger.info(
                    f"Pass 3 verification: "
                    f"rejection {verification['rejection_check']['covered']}/{verification['rejection_check']['total']}, "
                    f"evaluation {verification['evaluation_check']['covered']}/{verification['evaluation_check']['total']}, "
                    f"documents {verification['document_check']['covered']}/{verification['document_check']['total']}"
                )

                # ── Pass 3b: Auto-complete missing sections ──
                added = self._auto_complete_sections(structure, verification, analysis)
                if added > 0:
                    logger.info(f"Pass 3b: Auto-completed {added} missing sections")
                    _progress(f"🔧 自动补全 {added} 个缺失章节")
                    # Re-verify after completion
                    verification = self._verify_structure(analysis, structure)
                    structure["verification"] = verification
                    logger.info(
                        f"Pass 3b re-verification: "
                        f"evaluation {verification['evaluation_check']['covered']}/{verification['evaluation_check']['total']}"
                    )

                # ── Pass 3c: Inject material_refs programmatically ──
                refs_injected = self._inject_material_refs(structure, verification)
                if refs_injected:
                    logger.info(f"Pass 3c: Injected material_refs into {refs_injected} sections")

                # ── Inject linkage into each section for frontend ──
                linkage = verification.get("section_linkage", {})
                linked_count = 0
                for vol in structure.get("volumes", []):
                    for sec in vol.get("sections", []):
                        title = sec.get("title", "")
                        if title in linkage:
                            sec["linked_scoring"] = linkage[title].get("scoring_items", [])
                            sec["linked_rejection"] = linkage[title].get("rejection_items", [])
                            sec["linked_total_score"] = linkage[title].get("total_score", 0)
                            linked_count += 1
                        else:
                            sec["linked_scoring"] = []
                            sec["linked_rejection"] = []
                            sec["linked_total_score"] = 0
                if linked_count:
                    logger.info(f"Section linkage: {linked_count} sections linked to scoring/rejection items")
                    _progress(f"✅ Pass 3 完成: {linked_count} 个章节关联到评分/废标项")

                # ── Pass 3d: Auto-generate deviation/compliance tables ──
                deviation_sections = self._generate_deviation_tables(
                    analysis, verification, structure
                )
                if deviation_sections:
                    # Insert at the beginning of the first volume
                    first_vol = structure.get("volumes", [{}])[0]
                    existing_sections = first_vol.get("sections", [])
                    # Insert after form sections (投标函, 授权委托书 etc.)
                    # but before content sections
                    insert_idx = 0
                    for i, sec in enumerate(existing_sections):
                        if sec.get("type") in ("form",):
                            insert_idx = i + 1
                        else:
                            break
                    for j, dev_sec in enumerate(deviation_sections):
                        existing_sections.insert(insert_idx + j, dev_sec)
                    # Re-number all sections
                    for i, sec in enumerate(existing_sections, 1):
                        sec["order"] = i
                    first_vol["sections"] = existing_sections

                    # Rebuild order map with correct post-insertion numbers
                    # and regenerate deviation table content
                    new_order_map = {}
                    for sec in existing_sections:
                        t = sec.get("title", "")
                        o = sec.get("order", 0)
                        new_order_map[t] = f"第{o}章 {t}"
                    for dev_sec in deviation_sections:
                        self._patch_deviation_content(
                            dev_sec, new_order_map, verification
                        )

                    logger.info(
                        f"Pass 3d: Generated {len(deviation_sections)} "
                        f"deviation table(s), inserted at position {insert_idx + 1}"
                    )
                    _progress(f"📊 自动生成 {len(deviation_sections)} 个评审偏离表")

                # Store lot info in result
                if lot_info:
                    structure["lot_info"] = lot_info

                return structure

        except Exception as e:
            logger.error(f"Pass 2 structure generation failed: {e}")

        # Fallback: build structure from Pass 1 analysis manually
        logger.warning("Pass 2 failed, building structure from Pass 1 analysis")
        _progress("⚠️ Pass 2 大纲生成失败，降级到 Pass 1 构建模式...")
        fallback_structure = self._build_from_analysis(analysis)

        # Still run Pass 3 verification to get section_linkage (scoring/rejection)
        if fallback_structure and fallback_structure.get("volumes"):
            _progress("🔍 Pass 3: 对降级大纲执行废标/评分联动校验...")
            fallback_structure["tender_analysis"] = {
                "rejection_conditions": analysis.get("rejection_conditions", []),
                "evaluation_criteria": analysis.get("evaluation_criteria", []),
            }
            verification = self._verify_structure(analysis, fallback_structure)
            linkage = verification.get("section_linkage", {})
            linked_count = 0
            for vol in fallback_structure.get("volumes", []):
                for sec in vol.get("sections", []):
                    title = sec.get("title", "")
                    if title in linkage:
                        sec["linked_scoring"] = linkage[title].get("scoring_items", [])
                        sec["linked_rejection"] = linkage[title].get("rejection_items", [])
                        linked_count += 1
                    else:
                        sec["linked_scoring"] = []
                        sec["linked_rejection"] = []
            if linked_count:
                logger.info(f"Fallback section linkage: {linked_count} sections linked")
                _progress(f"✅ 降级模式联动完成: {linked_count} 个章节关联到评分/废标项")

        return fallback_structure

    @staticmethod
    def _extract_format_spec(raw_text, sections, file_path=''):
        """Extract bid document format specification from tender Chapter 6.

        Scans the tender for a '投标文件格式' section, then extracts:
        - Attachment list (附件1, 附件2, etc.) with names
        - Font specifications (from the docx file if available)
        - Table templates as XML files for exact cloning

        Returns a format_spec dict, or empty dict if no format chapter found.
        """
        import re
        import os
        from copy import deepcopy

        # ── Step 1: Find format chapter in text ──
        format_keywords = ['投标文件格式', '投标文件编制格式', '投标文件组成']
        has_format = False
        format_text = ''

        for kw in format_keywords:
            idx = raw_text.find(kw)
            if idx >= 0:
                has_format = True
                format_text = raw_text[idx:]
                break

        if not has_format:
            return {}

        # ── Step 2: Extract attachment list from text ──
        attachment_pattern = re.compile(
            r'附件\s*(\d+)\s*[：:]\s*(.+?)(?:\n|$)'
        )
        attachments = []
        seen_ids = set()
        for m in attachment_pattern.finditer(format_text):
            att_id = int(m.group(1))
            att_name = m.group(2).strip()
            # Clean trailing text from title
            # e.g. "投标一览表中内容进行报价；" → "投标一览表"
            # but keep "法定代表人（单位负责人）授权书" intact
            # Always strip '中内容' — definitive trailing text marker
            mid_idx = att_name.find('中内容')
            if mid_idx > 2:
                att_name = att_name[:mid_idx].strip()
            elif len(att_name) > 15:
                for sep in ['）。', '）；', '。', '；', ',', '，']:
                    idx2 = att_name.find(sep)
                    if idx2 > 2:
                        att_name = att_name[:idx2].strip()
                        break
                else:
                    # Handle （ only if no matching ）
                    pi = att_name.find('（')
                    if pi > 2:
                        ci = att_name.find('）', pi)
                        if ci < 0:
                            att_name = att_name[:pi].strip()
            if att_id not in seen_ids:
                seen_ids.add(att_id)
                attachments.append({
                    'id': f'附件{att_id}',
                    'order': att_id,
                    'title': att_name,
                    'tables': [],  # will be filled with table file paths
                })

        if not file_path or not os.path.exists(file_path):
            return {
                'has_format_chapter': True,
                'section_numbering': '附件' if attachments else '章',
                'font': {'name': 'Arial', 'size': 12, 'title_size': 15},
                'table_font': {'name': 'Arial', 'size': 11},
                'attachments': attachments,
            }

        # ── Step 3: Walk through docx body to extract tables with context ──
        try:
            from docx import Document as _Doc
            from lxml import etree

            doc = _Doc(file_path)

            # Create directory for table templates
            base_dir = os.path.dirname(file_path)
            tpl_dir = os.path.join(base_dir, 'table_templates')
            os.makedirs(tpl_dir, exist_ok=True)

            # Walk body elements in order, tracking current attachment
            format_started = False
            current_attachment = None  # index into attachments list
            font_samples = {}  # font_name → count
            table_font_spec = {'name': 'Arial', 'size': 11}
            font_spec = {'name': 'Arial', 'size': 12, 'title_size': 15}
            table_count = 0

            for element in doc.element.body:
                # ── Paragraph: check for format chapter start / attachment heading ──
                if element.tag.endswith('}p'):
                    for p in doc.paragraphs:
                        if p._element is element:
                            text = p.text.strip()

                            # Detect format chapter start
                            if not format_started:
                                if any(kw in text for kw in format_keywords):
                                    format_started = True

                            if not format_started:
                                break

                            # Track font samples
                            if p.runs:
                                for r in p.runs:
                                    fn = r.font.name
                                    if fn:
                                        font_samples[fn] = font_samples.get(fn, 0) + 1

                            # Detect attachment heading: "附件N：xxx" or standalone title
                            att_match = re.match(r'附件\s*(\d+)', text)
                            if att_match:
                                att_id = int(att_match.group(1))
                                # Find matching attachment in our list
                                for i, att in enumerate(attachments):
                                    if att['order'] == att_id:
                                        current_attachment = i
                                        break

                            # Also detect standalone section titles
                            standalone_map = {
                                '拟派实施人员表': None,
                                '拟派人员资历表': None,
                                '投标人情况表': None,
                            }
                            for st_title in standalone_map:
                                if st_title in text and len(text) < 30:
                                    # Find parent attachment (附件10 usually)
                                    # Keep current_attachment
                                    pass
                            break

                # ── Table: save XML if in format chapter ──
                elif element.tag.endswith('}tbl') and format_started:
                    for table in doc.tables:
                        if table._element is element:
                            # Extract table metadata
                            rows = len(table.rows)
                            cols = len(table.columns)

                            # Get headers
                            headers = []
                            for ci in range(cols):
                                try:
                                    h = table.cell(0, ci).text.strip()
                                    headers.append(h)
                                except Exception:
                                    headers.append('')

                            # Get column widths
                            col_widths = []
                            for col in table.columns:
                                try:
                                    w = col.width
                                    if w:
                                        col_widths.append(w)
                                    else:
                                        col_widths.append(0)
                                except Exception:
                                    col_widths.append(0)

                            # Get font info from table
                            for row in table.rows[:2]:
                                for cell in row.cells:
                                    for p in cell.paragraphs:
                                        for r in p.runs:
                                            if r.font.name:
                                                table_font_spec['name'] = r.font.name
                                            if r.font.size:
                                                table_font_spec['size'] = r.font.size.pt
                                            break
                                        break
                                    break

                            # Save table XML to file
                            table_count += 1
                            att_label = attachments[current_attachment]['id'] if current_attachment is not None else f'unknown_{table_count}'
                            xml_filename = f'table_{att_label}_{table_count}.xml'
                            xml_path = os.path.join(tpl_dir, xml_filename)

                            # Serialize the table element XML
                            xml_bytes = etree.tostring(
                                table._element,
                                xml_declaration=False,
                                encoding='unicode',
                            )
                            with open(xml_path, 'w', encoding='utf-8') as f:
                                f.write(xml_bytes)

                            table_info = {
                                'xml_path': xml_path,
                                'headers': headers,
                                'col_widths': col_widths,
                                'rows': rows,
                                'cols': cols,
                            }

                            # Attach to current attachment
                            if current_attachment is not None:
                                attachments[current_attachment]['tables'].append(table_info)
                                logger.info(
                                    f"  Table template saved: {att_label} "
                                    f"({rows}x{cols}) → {xml_filename}"
                                )
                            else:
                                logger.info(
                                    f"  Table template saved (unmatched): "
                                    f"({rows}x{cols}) → {xml_filename}"
                                )
                            break

            # Determine most common font
            if font_samples:
                most_common = max(font_samples, key=font_samples.get)
                font_spec['name'] = most_common
                if table_font_spec['name'] == 'Arial':
                    table_font_spec['name'] = most_common

        except Exception as e:
            logger.warning(f"Failed to extract table templates from docx: {e}")
            import traceback
            traceback.print_exc()

        format_spec = {
            'has_format_chapter': True,
            'section_numbering': '附件' if attachments else '章',
            'font': font_spec,
            'table_font': table_font_spec,
            'attachments': attachments,
        }

        tables_saved = sum(len(a.get('tables', [])) for a in attachments)
        logger.info(
            f"Format spec extracted: {len(attachments)} attachments, "
            f"{tables_saved} table templates, "
            f"font={font_spec['name']}/{font_spec['size']}pt, "
            f"table_font={table_font_spec['name']}/{table_font_spec['size']}pt"
        )

        return format_spec

    def _verify_structure(self, analysis, structure):
        # type: (Dict, Dict) -> Dict
        """Pass 3: Deterministic cross-reference verification.

        Compare Pass 1 analysis results against Pass 2 structure to check
        if all rejection conditions, evaluation criteria, and required
        documents are covered by the bid document structure.
        """
        # Collect all section titles for matching
        all_titles = []
        for vol in structure.get("volumes", []):
            for sec in vol.get("sections", []):
                all_titles.append(sec.get("title", ""))
        titles_text = " ".join(all_titles)

        def _find_matching_section(keywords):
            """Find a section title that matches any of the keywords."""
            if isinstance(keywords, str):
                keywords = [keywords]
            for kw in keywords:
                if not kw:
                    continue
                kw_lower = kw.lower().strip()
                for title in all_titles:
                    if kw_lower in title.lower():
                        return title
                # Fuzzy: check if key terms from kw appear in any title
                key_terms = [t for t in kw_lower.replace("（", " ").replace("）", " ").split() if len(t) >= 2]
                for title in all_titles:
                    title_lower = title.lower()
                    if any(term in title_lower for term in key_terms if len(term) >= 2):
                        return title
            return None

        # Collect rejection_risk sections
        rej_risk_titles = []
        for vol in structure.get("volumes", []):
            for sec in vol.get("sections", []):
                if sec.get("rejection_risk"):
                    rej_risk_titles.append(sec.get("title", ""))

        # Common keyword mapping: abstract condition → keywords to search in titles
        _CONDITION_KEYWORD_MAP = {
            "资格": ["资格", "资质", "证明", "营业执照", "执业"],
            "装订": ["装订", "格式", "投标文件"],
            "密封": ["密封", "投标文件"],
            "保证金": ["保证金", "担保"],
            "签署": ["签署", "签字", "公章", "投标函", "授权"],
            "完整": ["投标函", "开标", "报价"],
            "失信": ["信用", "失信", "声明"],
            "有效期": ["有效期", "投标函"],
            "报价": ["报价", "开标一览表", "价格"],
        }

        def _match_condition(condition_text, related_doc):
            """Try to match a rejection condition to a section via multiple strategies."""
            # Strategy 1: Direct match on related_document
            if related_doc:
                m = _find_matching_section([related_doc])
                if m:
                    return m

            # Strategy 2: Keyword map
            for trigger, search_terms in _CONDITION_KEYWORD_MAP.items():
                if trigger in condition_text:
                    for term in search_terms:
                        for title in all_titles:
                            if term in title:
                                return title

            # Strategy 3: Check if any rejection_risk section relates
            cond_lower = condition_text.lower()
            for title in rej_risk_titles:
                title_lower = title.lower()
                # If condition and title share meaningful Chinese chars
                shared = sum(1 for c in title_lower if c in cond_lower and len(c.encode('utf-8')) > 1)
                if shared >= 2:
                    return title

            # Strategy 4: Fall back to regular fuzzy
            return _find_matching_section([condition_text])

        # ── Rejection conditions check ──
        rejection_items = []
        for rc in analysis.get("rejection_conditions", []):
            condition = rc.get("condition", "")
            related = rc.get("related_document", "")
            matched = _match_condition(condition, related)
            rejection_items.append({
                "condition": condition,
                "related_document": related,
                "source": rc.get("source", ""),
                "status": "covered" if matched else "missing",
                "matched_section": matched,
            })

        rej_covered = sum(1 for r in rejection_items if r["status"] == "covered")

        # ── Evaluation criteria check ──
        # Synonym mapping for common evaluation terms → section titles
        _EVAL_SYNONYM_MAP = {
            '技术方案': ['技术', '方案', '实施', '解决方案', '技术路线'],
            '服务方案': ['服务', '运维', '售后', '服务保障', '服务承诺', '服务方案'],
            '团队': ['人员', '律师', '团队', '简历', '拟投入', '拟委派', '项目经理', '配置', '成员'],
            '人员': ['人员', '团队', '配置', '律师', '成员', '项目组'],
            '人员构成': ['人员', '团队', '配置', '律师', '成员', '项目组'],
            '业绩': ['业绩', '案例', '项目经验', '类似项目', '合同'],
            '报价': ['报价', '价格', '费用', '开标', '一览表', '投标报价'],
            '资质': ['资质', '证书', '营业执照', '许可证', '认证'],
            '管理': ['管理', '质量', '进度', '风控', '安全', '保密'],
            '培训': ['培训', '知识转移', '交接'],
            '应急': ['应急', '预案', '备份', '容灾'],
            # New synonyms for previously-unmatched items
            '处罚': ['合规', '处罚', '声明', '信用', '诚信'],
            '合规': ['合规', '处罚', '声明', '信用'],
            '响应': ['响应', '投标文件', '说明', '偏离'],
            '分所': ['分所', '覆盖', '介绍', '概况', '律所'],
            '综合': ['综合', '实力', '介绍', '概况'],
            '质量': ['质量', '控制', '管理', '保障'],
        }

        # Aggregate eval items: these are parent items whose score is
        # the sum of sub-items. If all sub-topics are covered, the
        # aggregate is considered covered.
        _AGGREGATE_ITEMS = {
            '综合实力': ['业绩', '荣誉', '人员', '团队', '资质', '介绍', '分所'],
        }

        def _eval_match(item_name, needed):
            """Try to match an evaluation criterion to a section using synonyms."""
            # Priority 0: Deterministic mapping from EVAL_TO_SECTION_MAP
            for key, val in RequirementExtractionSkill.EVAL_TO_SECTION_MAP.items():
                if key in item_name:
                    if val is None:
                        return None  # aggregate item, skip
                    mapped_title = val[0]
                    # Check if this mapped title exists in actual sections
                    for title in all_titles:
                        if mapped_title in title or title in mapped_title:
                            return title
            # Priority 1: Direct match (title contains needed or item_name)
            matched = _find_matching_section([needed, item_name])
            if matched:
                return matched
            # Priority 2: Precise item_name containment
            for title in all_titles:
                if item_name and len(item_name) >= 2:
                    if item_name in title or title in item_name:
                        return title
            # Priority 3: Synonym-based matching (broadest)
            for trigger, synonyms in _EVAL_SYNONYM_MAP.items():
                if trigger in item_name or trigger in (needed or ''):
                    for syn in synonyms:
                        for title in all_titles:
                            if syn in title:
                                return title
            return None

        eval_items = []
        for ec in analysis.get("evaluation_criteria", []):
            item_name = ec.get("item", "")
            needed = ec.get("bid_section_needed", "")
            matched = _eval_match(item_name, needed)
            eval_items.append({
                "item": item_name,
                "max_score": ec.get("max_score", 0),
                "description": ec.get("description", ""),
                "material_evidence": ec.get("material_evidence", ""),
                "category": ec.get("category", ""),
                "bid_section_needed": needed,
                "sub_criteria": ec.get("sub_criteria", []),
                "status": "covered" if matched else "missing",
                "matched_section": matched,
            })

        # Post-process: mark aggregate items as covered if their sub-topics
        # are represented by other covered eval items or existing sections
        for ei in eval_items:
            if ei["status"] != "missing":
                continue
            item_name = ei["item"]
            for agg_name, sub_topics in _AGGREGATE_ITEMS.items():
                if agg_name in item_name:
                    # Check if sub-topics are covered by existing sections
                    sub_covered = 0
                    for topic in sub_topics:
                        for title in all_titles:
                            if topic in title:
                                sub_covered += 1
                                break
                    if sub_covered >= 2:
                        # At least 2 sub-topics have corresponding sections
                        ei["status"] = "covered"
                        ei["matched_section"] = f"(聚合项: {sub_covered}/{len(sub_topics)}子项已覆盖)"
                        logger.info(
                            f"  Aggregate eval '{item_name}': "
                            f"{sub_covered}/{len(sub_topics)} sub-topics covered by sections"
                        )
                    break

        eval_covered = sum(1 for e in eval_items if e["status"] == "covered")
        # Log per-item coverage for debugging
        for e in eval_items:
            status_icon = "✅" if e["status"] == "covered" else "❌"
            logger.info(
                f"  Eval item {status_icon} [{e.get('max_score', 0)}分] "
                f"'{e['item']}' → {e.get('matched_section', 'NO MATCH')}"
            )

        # ── Required documents check ──
        doc_items = []
        for doc in analysis.get("bid_composition", {}).get("required_documents", []):
            doc_name = doc.get("name", "")
            matched = _find_matching_section([doc_name])
            doc_items.append({
                "name": doc_name,
                "category": doc.get("category", ""),
                "is_mandatory": doc.get("is_mandatory", True),
                "source": doc.get("source", ""),
                "status": "covered" if matched else "missing",
                "matched_section": matched,
            })

        doc_covered = sum(1 for d in doc_items if d["status"] == "covered")

        # ── Format & deadline info ──
        fmt = analysis.get("format_requirements", {})
        deadline = analysis.get("deadline_info", {})

        # ── Build forward index: section → linked scoring + rejection ──
        section_linkage = {}  # type: Dict[str, Dict]
        for ei in eval_items:
            matched = ei.get("matched_section", "")
            if matched:
                if matched not in section_linkage:
                    section_linkage[matched] = {
                        "scoring_items": [], "rejection_items": [],
                        "total_score": 0,
                    }
                def _safe_score(v):
                    try:
                        return int(v) if v else 0
                    except (ValueError, TypeError):
                        return 0

                section_linkage[matched]["scoring_items"].append({
                    "item": ei["item"],
                    "max_score": _safe_score(ei.get("max_score", 0)),
                    "description": ei.get("description", ""),
                    "sub_criteria": ei.get("sub_criteria", []),
                })
                section_linkage[matched]["total_score"] += _safe_score(ei.get("max_score", 0))

        for ri in rejection_items:
            matched = ri.get("matched_section", "")
            if matched:
                if matched not in section_linkage:
                    section_linkage[matched] = {
                        "scoring_items": [], "rejection_items": [],
                        "total_score": 0,
                    }
                section_linkage[matched]["rejection_items"].append({
                    "condition": ri["condition"],
                    "source": ri.get("source", ""),
                })

        # ── Validate sub_criteria score sums ──
        score_warnings = []
        for ei in eval_items:
            subs = ei.get("sub_criteria", [])
            if subs:
                sub_total = sum(s.get("score", 0) for s in subs)
                max_score = ei.get("max_score", 0)
                if max_score > 0 and sub_total != max_score:
                    score_warnings.append({
                        "item": ei["item"],
                        "max_score": max_score,
                        "sub_total": sub_total,
                        "diff": max_score - sub_total,
                    })
        if score_warnings:
            logger.warning(f"Score sum mismatch in {len(score_warnings)} items: "
                          f"{score_warnings}")

        return {
            "rejection_check": {
                "total": len(rejection_items),
                "covered": rej_covered,
                "items": rejection_items,
            },
            "evaluation_check": {
                "total": len(eval_items),
                "covered": eval_covered,
                "items": eval_items,
            },
            "document_check": {
                "total": len(doc_items),
                "covered": doc_covered,
                "items": doc_items,
            },
            "format_info": fmt,
            "deadline_info": deadline,
            "section_linkage": section_linkage,
            "score_warnings": score_warnings,
        }

    def _build_from_analysis(self, analysis):
        # type: (Dict) -> Dict
        """Fallback: build bid structure directly from Pass 1 results."""
        sections = []
        order = 1

        # Add required documents from bid_composition
        for doc in analysis.get("bid_composition", {}).get("required_documents", []):
            sections.append({
                "order": order,
                "title": doc.get("name", f"文件{order}"),
                "type": doc.get("category", "narrative"),
                "required": doc.get("is_mandatory", True),
                "rejection_risk": doc.get("is_mandatory", False),
                "score_weight": 0,
                "content_hints": doc.get("source", ""),
                "content_outline": [],
                "material_refs": [],
                "data_fields": [],
            })
            order += 1

        # Add sections for evaluation criteria that don't have matching documents
        existing_titles = {s["title"] for s in sections}
        for crit in analysis.get("evaluation_criteria", []):
            needed = crit.get("bid_section_needed", "")
            if needed and needed not in existing_titles:
                sections.append({
                    "order": order,
                    "title": needed,
                    "type": "narrative",
                    "required": True,
                    "rejection_risk": False,
                    "score_weight": crit.get("max_score", 0),
                    "content_hints": crit.get("description", ""),
                    "content_outline": [crit.get("description", "")] if crit.get("description") else [],
                    "material_refs": [],
                    "data_fields": [],
                })
                existing_titles.add(needed)
                order += 1

        return {
            "bid_title": analysis.get("project_info", {}).get("project_name", "投标文件") or "投标文件",
            "volumes": [{"name": "投标文件", "sections": sections}],
            "rejection_items": [
                {
                    "description": r.get("condition", ""),
                    "related_sections": [r.get("related_document", "")],
                    "severity": "critical",
                }
                for r in analysis.get("rejection_conditions", [])
            ],
            "tender_analysis": {
                "project_info": analysis.get("project_info", {}),
                "rejection_conditions": analysis.get("rejection_conditions", []),
                "evaluation_criteria": analysis.get("evaluation_criteria", []),
                "qualification_requirements": analysis.get("qualification_requirements", []),
                "format_requirements": analysis.get("format_requirements", {}),
                "deadline_info": analysis.get("deadline_info", {}),
            },
            "qualification_requirements": analysis.get("qualification_requirements", []),
            "format_requirements": analysis.get("format_requirements", {}),
            "evaluation_criteria": analysis.get("evaluation_criteria", []),
            "deadline_info": analysis.get("deadline_info", {}),
        }

    # ── V2: Section title classification (kept as fallback) ──

    async def _structure_from_sections(self, llm, sections, raw_text):
        # type: (Any, List[Dict], str) -> Dict
        """V2: Build structure from parsed section titles, LLM classifies."""
        logger.info(f"V2 structure mode: {len(sections)} raw sections")

        sections = self._filter_bid_sections(sections)

        section_lines = []
        for i, sec in enumerate(sections, 1):
            title = sec.get("title", f"第{i}节")
            section_lines.append(f"{i}. {title}")

        section_list_text = "\n".join(section_lines)
        tender_context = raw_text[:8000] if raw_text else "暂无原文"

        prompt = BATCH_CLASSIFY_PROMPT.format(
            tender_context=tender_context,
            section_list=section_list_text,
        )

        try:
            response = await llm.generate(prompt, system=BATCH_CLASSIFY_SYSTEM)
            classifications = _safe_parse_json(response)
        except Exception as e:
            logger.error(f"LLM classification failed: {e}")
            classifications = None

        classify_map = {}  # type: Dict[str, Dict]
        if classifications and isinstance(classifications, list):
            for item in classifications:
                title = item.get("title", "")
                if title:
                    classify_map[title] = {
                        "type": item.get("type", "narrative"),
                        "content_hints": item.get("content_hints", ""),
                        "data_fields": item.get("data_fields", []),
                    }

        bid_sections = []
        for i, sec in enumerate(sections):
            title = sec.get("title", f"第{i+1}节")
            original_type = sec.get("section_type", "narrative")
            content = sec.get("content", "")
            classification = classify_map.get(title, {})
            sec_type = classification.get("type", original_type)
            hints = classification.get("content_hints", content[:100] if content else "")
            fields = classification.get("data_fields", [])

            bid_sections.append({
                "order": i + 1,
                "title": title,
                "type": sec_type,
                "required": True,
                "rejection_risk": False,
                "content_hints": hints,
                "data_fields": fields,
            })

        return {
            "bid_title": "投标文件",
            "volumes": [{"name": "投标文件", "sections": bid_sections}],
            "qualification_requirements": [],
            "format_requirements": {},
            "evaluation_criteria": [],
            "deadline_info": {},
        }

    def _filter_bid_sections(self, sections):
        # type: (List[Dict]) -> List[Dict]
        """Filter tender_parsing sections to keep only bid-relevant headings."""
        import re

        SKIP_PATTERNS = [
            r'投标人应当', r'投标人递交', r'投标人没有',
            r'招标人有权', r'招标人不予',
            r'投标文件应当使用不褪色', r'应当按照招标文件',
            r'应当认真阅读', r'并加盖单位公章',
        ]
        skip_regex = re.compile('|'.join(SKIP_PATTERNS))

        MAX_TITLE_LEN = 60
        seen_titles = set()  # type: set
        filtered = []

        for sec in sections:
            title = sec.get("title", "").strip()
            if not title or len(title) > MAX_TITLE_LEN:
                continue
            if skip_regex.search(title):
                continue
            if title in seen_titles:
                continue
            seen_titles.add(title)
            filtered.append(sec)

        logger.info(f"Section filter: {len(sections)} → {len(filtered)}")
        return filtered

    # ── V1: Legacy single-pass extraction ──

    async def _extract_full(self, llm, raw_text):
        # type: (Any, str) -> Any
        """Legacy: Single pass extraction."""
        EXTRACTION_PROMPT = """请分析以下招标文件内容，提取投标文件需要包含的所有章节和要求。

【招标文件内容】
{tender_text}

请输出 JSON 结构（含 bid_title, volumes, qualification_requirements 等），只输出 JSON。"""

        prompt = EXTRACTION_PROMPT.format(tender_text=raw_text[:self.MAX_CHUNK_SIZE])
        response = await llm.generate(prompt, system=ANALYSIS_SYSTEM)
        return _safe_parse_json(response)

    def _fallback_from_sections(self, sections):
        # type: (List[Dict]) -> Dict
        """Create basic structure from parsed sections when all else fails."""
        logger.warning("Using fallback structure from parsed sections")
        bid_sections = []
        for i, sec in enumerate(sections):
            bid_sections.append({
                "order": i + 1,
                "title": sec.get("title", f"第{i+1}节"),
                "type": sec.get("section_type", "narrative"),
                "required": True,
                "rejection_risk": False,
                "content_hints": sec.get("content", "")[:100],
                "content_outline": [],
                "material_refs": [],
                "data_fields": [],
            })

        return {
            "bid_title": "投标文件",
            "volumes": [{"name": "投标文件", "sections": bid_sections}],
            "qualification_requirements": [],
            "format_requirements": {},
            "evaluation_criteria": [],
            "deadline_info": {},
        }

    def _refine_section_types(self, result):
        # type: (Dict) -> Dict
        """Use local BGE classifier to correct section types if available."""
        try:
            from app.core.rag.section_classifier import section_classifier

            corrections = 0
            for volume in result.get("volumes", []):
                titles = [s.get("title", "") for s in volume.get("sections", [])]
                if not titles:
                    continue

                classifications = section_classifier.classify_batch(titles)

                for section, (predicted_type, confidence) in zip(
                    volume.get("sections", []), classifications
                ):
                    llm_type = section.get("type", "narrative")
                    if predicted_type != llm_type and confidence > 0.7:
                        logger.debug(
                            f"Type correction: '{section.get('title')}' "
                            f"{llm_type} → {predicted_type} (conf={confidence:.2f})"
                        )
                        section["type"] = predicted_type
                        section["type_confidence"] = round(confidence, 3)
                        section["type_source"] = "classifier"
                        corrections += 1
                    else:
                        section["type_source"] = "llm"

            if corrections > 0:
                logger.info(f"Section type classifier corrected {corrections} types")

        except Exception as e:
            logger.warning(f"Section classifier not available: {e}")

        return result

    # ── Evaluation item → section deterministic mapping ──
    EVAL_TO_SECTION_MAP = {
        # eval_keyword → (section_title, section_type, material_refs)
        "业绩": ("律所业绩", "narrative", ["近年类似业绩项目"]),
        "律所业绩": ("律所业绩", "narrative", ["近年类似业绩项目"]),
        "项目经验": ("律所业绩", "narrative", ["近年类似业绩项目"]),
        "团队": ("项目团队配置", "narrative", ["项目经理及核心成员简历"]),
        "人员": ("项目团队配置", "narrative", ["项目经理及核心成员简历"]),
        "人员构成": ("项目团队配置", "narrative", ["项目经理及核心成员简历"]),
        "资质": ("资格审查资料", "qualification", ["营业执照", "执业许可证"]),
        "荣誉": ("荣誉奖项", "qualification", ["行业排名证明", "获奖证书"]),
        "荣誉奖项": ("荣誉奖项", "qualification", ["行业排名证明", "获奖证书"]),
        "方案": ("服务方案", "narrative", []),
        "服务方案": ("服务方案", "narrative", []),
        "服务方案编制": ("服务方案", "narrative", []),
        "质量": ("服务质量控制方案", "narrative", []),
        "服务质量": ("服务质量控制方案", "narrative", []),
        "服务质量控制": ("服务质量控制方案", "narrative", []),
        "报价": ("报价文件", "form", []),
        # "综合实力" is an aggregate eval item — map sub-items, not itself.
        # When LLM produces sub_criteria like "分所覆盖" under "综合实力",
        # those sub-items should match to existing sections (律所荣誉, etc).
        "综合实力": None,  # Skip — it's an aggregate, handled by sub-criteria
        "分所": None,       # Skip — covered by 律所介绍/概况
        "分所覆盖": None,   # Skip — covered by 律所介绍/概况
        "处罚": ("合规声明", "form", []),
        "处罚情况": ("合规声明", "form", []),
        "投标文件响应": ("投标文件响应说明", "narrative", []),
        "财务": ("财务状况", "qualification", ["审计报告"]),
    }

    # ── Topic groups for semantic overlap detection ──
    # If ANY keyword from a group appears in an existing section title,
    # the whole group is considered "covered".
    _TOPIC_GROUPS = [
        # (group_name, keywords that indicate this topic)
        ("业绩", ["业绩", "项目经验", "类似项目", "案例", "代表项目", "服务案例"]),
        ("团队", ["团队", "人员", "简历", "律师", "拟投入", "成员", "项目组"]),
        ("荣誉", ["荣誉", "奖项", "排名", "评级", "获奖"]),
        ("方案", ["服务方案", "实施方案", "技术方案", "工作方案"]),
        ("质量", ["质量控制", "质量管理", "质量保证"]),
        ("资质", ["资格审查", "资质", "执照", "许可"]),
        ("合规", ["合规", "处罚", "诚信", "声明"]),
        ("综合实力", ["综合实力", "律所介绍", "公司简介", "律所概况"]),
        ("商务", ["商务部分", "商务文件"]),
    ]

    # ── Banned section names (must NOT be auto-created) ──
    _BANNED_SECTION_NAMES = [
        "商务部分", "综合实力", "律所综合实力",
        "技术部分", "其他文件和资料", "补充材料",
        "附件", "其他资料",
    ]

    # ── Keywords → material_refs injection rules ──
    _MATERIAL_REF_RULES = [
        # (title_keywords, material_refs_to_inject)
        (["团队", "人员", "简历", "律师"], ["项目经理及核心成员简历"]),
        (["业绩", "项目经验", "案例", "合同"], ["近年类似业绩项目"]),
        (["资格", "资质", "证书", "执照", "许可"], ["资质证书"]),
        (["荣誉", "奖项", "排名", "评级"], ["行业排名证明", "获奖证书"]),
        (["财务", "审计", "报表"], ["审计报告"]),
    ]

    def _auto_complete_sections(self, structure, verification, analysis):
        # type: (Dict, Dict, Dict) -> int
        """Auto-complete: add missing sections for uncovered evaluation
        criteria and required documents. Returns number of sections added."""
        added = 0

        # Collect existing section titles
        existing_titles = set()
        target_volume = None
        for vol in structure.get("volumes", []):
            for sec in vol.get("sections", []):
                existing_titles.add(sec.get("title", ""))
            if target_volume is None:
                target_volume = vol

        if target_volume is None:
            return 0

        sections = target_volume.get("sections", [])
        max_order = max((s.get("order", 0) for s in sections), default=0)

        # ── Step 1: Auto-complete from uncovered evaluation criteria ──
        for ei in verification.get("evaluation_check", {}).get("items", []):
            if ei.get("status") == "covered":
                continue

            item_name = ei.get("item", "")
            bid_section_needed = ei.get("bid_section_needed", "")

            # Try to find a mapping
            sec_title = None
            sec_type = "narrative"
            sec_refs = []

            # Strategy 1: Direct mapping from EVAL_TO_SECTION_MAP
            for key, val in self.EVAL_TO_SECTION_MAP.items():
                if key in item_name:
                    if val is None:
                        # Explicitly skipped — aggregate or covered elsewhere
                        sec_title = None
                        logger.debug(
                            f"  Skip eval item '{item_name}': mapped to None (aggregate)"
                        )
                        break
                    sec_title, sec_type, sec_refs = val
                    break

            # If explicitly skipped (None mapping), skip entirely
            if sec_title is None and any(k in item_name for k, v in self.EVAL_TO_SECTION_MAP.items() if v is None):
                continue

            # Strategy 2: Use bid_section_needed from Pass 1
            if not sec_title and bid_section_needed:
                sec_title = bid_section_needed

            # Strategy 3: Use item_name as title
            if not sec_title:
                sec_title = item_name

            # ── Banned name filter ──
            if any(ban in sec_title for ban in self._BANNED_SECTION_NAMES):
                logger.info(
                    f"  Skip auto-add '{sec_title}': matches banned pattern"
                )
                continue

            # Skip if a section with similar title already exists
            if sec_title in existing_titles:
                continue

            # ── Topic-group overlap detection ──
            # Check if the new section's topic is already covered by
            # an existing section, using _TOPIC_GROUPS for semantic matching.
            _topic_overlap = False

            # Step A: Find which topic groups the NEW section belongs to
            new_topic_groups = set()
            for group_name, keywords in self._TOPIC_GROUPS:
                for kw in keywords:
                    if kw in sec_title or kw in item_name:
                        new_topic_groups.add(group_name)
                        break

            # Step B: Find which topic groups the EXISTING sections cover
            existing_topic_groups = set()
            for existing_title in existing_titles:
                if not existing_title:
                    continue
                for group_name, keywords in self._TOPIC_GROUPS:
                    for kw in keywords:
                        if kw in existing_title:
                            existing_topic_groups.add(group_name)
                            break

            # Step C: If any topic group overlaps, skip
            overlap_groups = new_topic_groups & existing_topic_groups
            if overlap_groups:
                _topic_overlap = True
                logger.debug(
                    f"  Skip auto-add '{sec_title}': topic groups "
                    f"{overlap_groups} already covered by existing sections"
                )

            # Step D: Direct containment check (fallback)
            if not _topic_overlap:
                for existing_title in existing_titles:
                    if not existing_title:
                        continue
                    if sec_title in existing_title or existing_title in sec_title:
                        _topic_overlap = True
                        break

            if _topic_overlap:
                logger.debug(
                    f"  Skip auto-add '{sec_title}': overlaps with existing section"
                )
                continue

            max_order += 1
            new_section = {
                "order": max_order,
                "title": sec_title,
                "type": sec_type,
                "required": True,
                "rejection_risk": False,
                "score_weight": ei.get("max_score", 0),
                "content_hints": ei.get("description", ""),
                "content_outline": [
                    f"{sub.get('item', '')} ← {sub.get('description', '')}"
                    for sub in ei.get("sub_criteria", [])
                ] if ei.get("sub_criteria") else [ei.get("description", "")],
                "material_refs": sec_refs,
                "data_fields": [],
                "source_reference": f"评分项自动补全: {item_name}",
                "_auto_completed": True,
            }
            sections.append(new_section)
            existing_titles.add(sec_title)
            added += 1
            logger.info(
                f"  Auto-added section: [{sec_type}] {sec_title} "
                f"(eval: {item_name}, {ei.get('max_score', 0)}分)"
            )

        # ── Step 2: Auto-complete from uncovered required documents ──
        for di in verification.get("document_check", {}).get("items", []):
            if di.get("status") == "covered":
                continue
            if not di.get("is_mandatory", True):
                continue

            doc_name = di.get("name", "")
            if not doc_name or doc_name in existing_titles:
                continue
            # ── Banned name filter for documents too ──
            if any(ban in doc_name for ban in self._BANNED_SECTION_NAMES):
                logger.info(
                    f"  Skip auto-add document '{doc_name}': matches banned pattern"
                )
                continue

            max_order += 1
            doc_type = di.get("category", "form")
            if doc_type not in ("form", "qualification", "narrative", "table"):
                doc_type = "form"

            new_section = {
                "order": max_order,
                "title": doc_name,
                "type": doc_type,
                "required": True,
                "rejection_risk": True,  # mandatory document → rejection risk
                "score_weight": 0,
                "content_hints": f"必须提供: {doc_name}",
                "content_outline": [],
                "material_refs": [],
                "data_fields": [],
                "source_reference": f"必须文件自动补全",
                "_auto_completed": True,
            }
            sections.append(new_section)
            existing_titles.add(doc_name)
            added += 1
            logger.info(f"  Auto-added required document: [{doc_type}] {doc_name}")

        return added

    def _generate_deviation_tables(self, analysis, verification, structure):
        # type: (Dict, Dict, Dict) -> List[Dict]
        """Auto-generate deviation/compliance tables from eval linkage.

        Produces one section per scoring category (商务/技术/价格).
        Each section contains a markdown table mapping evaluation items
        to their corresponding bid chapter locations.
        """
        eval_items = verification.get("evaluation_check", {}).get("items", [])
        if not eval_items:
            return []

        # Build section order lookup: title → "第X章"
        section_order_map = {}  # type: Dict[str, str]
        for vol in structure.get("volumes", []):
            for sec in vol.get("sections", []):
                order = sec.get("order", 0)
                title = sec.get("title", "")
                section_order_map[title] = f"第{order}章 {title}"

        # Group eval items by category
        categories = {}  # type: Dict[str, List]
        for ei in eval_items:
            cat = ei.get("category", "")
            # Infer category from item content if Qwen didn't provide it
            if not cat:
                item_name = ei.get("item", "")
                # Heuristic: 报价/价格 → 价格; 方案/质量/人员 → 技术; else → 商务
                if any(kw in item_name for kw in ["报价", "价格", "费用"]):
                    cat = "价格"
                elif any(kw in item_name for kw in [
                    "方案", "质量", "人员", "团队", "业绩", "服务"
                ]):
                    cat = "技术"
                else:
                    cat = "商务"
            if cat not in categories:
                categories[cat] = []
            categories[cat].append(ei)

        # Desired order for tables
        cat_order = ["商务", "技术", "价格"]
        result_sections = []

        for cat in cat_order:
            items = categories.get(cat, [])
            if not items:
                continue

            # Build markdown table
            lines = [
                f"## {cat}评分偏离表\n",
                "| 序号 | 评分项目 | 分值 | 招标文件评分要求 | 材料依据 | 对应投标文件位置 | 偏离说明 | 备注 |",
                "|------|---------|------|----------------|----------|----------------|---------|------|",
            ]
            for idx, ei in enumerate(items, 1):
                item_name = ei.get("item", "—")
                score = ei.get("max_score", 0)
                desc = ei.get("description", "—")
                # Truncate long descriptions for table readability
                if len(desc) > 50:
                    desc = desc[:47] + "..."
                # Material evidence from scoring table
                evidence = ei.get("material_evidence", "")
                if evidence and len(evidence) > 40:
                    evidence = evidence[:37] + "..."
                if not evidence:
                    evidence = "—"
                matched = ei.get("matched_section", "")
                if matched and not matched.startswith("("):
                    location = section_order_map.get(matched, matched)
                else:
                    location = "详见正文"
                status = "covered" if ei.get("status") == "covered" else "missing"
                deviation = "无偏离" if status == "covered" else "待补充"
                lines.append(
                    f"| {idx} | {item_name} | {score} | {desc} "
                    f"| {evidence} | {location} | {deviation} | |"
                )

            content = "\n".join(lines)

            result_sections.append({
                "title": f"{cat}评分偏离表",
                "type": "table",
                "content_hints": f"根据招标文件{cat}评分表自动生成的偏离/响应索引表",
                "data_fields": [],
                "content_outline": [],
                "material_refs": [],
                "rejection_risk": False,
                "linked_scoring": [],
                "linked_rejection": [],
                "linked_total_score": 0,
                # Pre-fill content so content_generation skips LLM
                "_deviation_table_content": content,
            })

        if result_sections:
            logger.info(
                f"Generated deviation tables: "
                f"{', '.join(f'{cat}({len(items)}项)' for cat, items in categories.items() if items)}"
            )

        return result_sections

    def _patch_deviation_content(self, dev_sec, order_map, verification):
        # type: (Dict, Dict[str, str], Dict) -> None
        """Regenerate deviation table content with correct chapter numbers.

        Called after deviation sections are inserted and all sections
        re-numbered, so the order_map reflects final positions.
        """
        title = dev_sec.get("title", "")
        # Determine which category this table is for
        cat = title.replace("评分偏离表", "")  # "商务" / "技术" / "价格"

        # Filter eval items for this category
        eval_items = verification.get("evaluation_check", {}).get("items", [])
        cat_items = []
        for ei in eval_items:
            ei_cat = ei.get("category", "")
            if not ei_cat:
                item_name = ei.get("item", "")
                if any(kw in item_name for kw in ["报价", "价格", "费用"]):
                    ei_cat = "价格"
                elif any(kw in item_name for kw in [
                    "方案", "质量", "人员", "团队", "业绩", "服务"
                ]):
                    ei_cat = "技术"
                else:
                    ei_cat = "商务"
            if ei_cat == cat:
                cat_items.append(ei)

        if not cat_items:
            return

        lines = [
            f"## {cat}评分偏离表\n",
            "| 序号 | 评分项目 | 分值 | 招标文件评分要求 | 材料依据 | 对应投标文件位置 | 偏离说明 | 备注 |",
            "|------|---------|------|----------------|----------|----------------|---------|------|",
        ]
        for idx, ei in enumerate(cat_items, 1):
            item_name = ei.get("item", "—")
            score = ei.get("max_score", 0)
            desc = ei.get("description", "—")
            if len(desc) > 50:
                desc = desc[:47] + "..."
            evidence = ei.get("material_evidence", "")
            if evidence and len(evidence) > 40:
                evidence = evidence[:37] + "..."
            if not evidence:
                evidence = "—"
            matched = ei.get("matched_section", "")
            if matched and not matched.startswith("("):
                location = order_map.get(matched, matched)
            else:
                location = "详见正文"
            status = "covered" if ei.get("status") == "covered" else "missing"
            deviation = "无偏离" if status == "covered" else "待补充"
            lines.append(
                f"| {idx} | {item_name} | {score} | {desc} "
                f"| {evidence} | {location} | {deviation} | |"
            )

        dev_sec["_deviation_table_content"] = "\n".join(lines)

    def _inject_material_refs(self, structure, verification):
        # type: (Dict, Dict) -> int
        """Programmatically inject material_refs into sections based on
        title keywords and linked scoring items. Returns count of sections updated."""
        injected = 0
        linkage = verification.get("section_linkage", {})

        for vol in structure.get("volumes", []):
            for sec in vol.get("sections", []):
                title = sec.get("title", "")
                existing_refs = sec.get("material_refs", [])

                new_refs = list(existing_refs)  # preserve any LLM-generated refs

                # Rule 1: Title keyword matching
                for keywords, refs_to_add in self._MATERIAL_REF_RULES:
                    if any(kw in title for kw in keywords):
                        for ref in refs_to_add:
                            if ref not in new_refs:
                                new_refs.append(ref)

                # Rule 2: Linked scoring items
                sec_linkage = linkage.get(title, {})
                for scoring_item in sec_linkage.get("scoring_items", []):
                    item_name = scoring_item.get("item", "")
                    for key, val in self.EVAL_TO_SECTION_MAP.items():
                        if val is None:
                            continue
                        _, _, refs = val
                        if key in item_name:
                            for ref in refs:
                                if ref not in new_refs:
                                    new_refs.append(ref)

                if len(new_refs) > len(existing_refs):
                    sec["material_refs"] = new_refs
                    injected += 1

        return injected
