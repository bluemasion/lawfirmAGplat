"""投标文件生成 API — 上传解析 + Qwen 生成标书框架，全链路打通。"""

from fastapi import APIRouter, UploadFile, File, Form
from fastapi.responses import StreamingResponse, FileResponse, HTMLResponse
from pydantic import BaseModel
from typing import Optional
import json, zipfile, io, re
import xml.etree.ElementTree as ET

from app.core.llm import get_llm
from app.utils.logger import logger

router = APIRouter()

# ─────────────────────────────────────────────
#  1. 上传解析端点：提取招标文件关键信息
# ─────────────────────────────────────────────

PARSE_SYSTEM_PROMPT = """你是一位专业的招标文件分析专家。给你一份招标文件的原文内容，你需要从中提取以下结构化信息并以 JSON 格式返回。

请严格按照以下 JSON 格式返回，不要返回其他任何内容：
{
  "project_name": "项目名称",
  "client_name": "招标人/采购人名称",
  "project_id": "项目编号",
  "budget": "预算金额/最高限价",
  "guarantee_amount": "投标保证金金额",
  "validity_days": "投标有效期(天数)",
  "bid_method": "招标方式(公开/邀请)",
  "evaluation_method": "评标方法(最低价/综合评分等)",
  "submission_deadline": "投标截止时间",
  "requirements": ["关键资质/技术要求1", "关键资质/技术要求2"],
  "key_terms": ["重要商务条款1", "重要商务条款2"],
  "disqualification_risks": ["否决条件1", "否决条件2"]
}

注意:
- 如果某个字段在文件中找不到，填写 "未明确" 
- 金额保留原文单位
- requirements 提取最关键的 3-5 条
- disqualification_risks 提取可能导致投标被否决的条件"""


def extract_text_from_docx(file_bytes: bytes) -> str:
    """从 docx 文件字节中提取纯文本（不依赖 python-docx）"""
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
            xml_content = zf.read("word/document.xml")
        
        root = ET.fromstring(xml_content)
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        
        paragraphs = []
        for p in root.findall('.//w:p', ns):
            texts = [t.text for t in p.findall('.//w:t', ns) if t.text]
            line = ''.join(texts).strip()
            if line:
                paragraphs.append(line)
        
        return '\n'.join(paragraphs)
    except Exception as e:
        logger.error(f"DOCX extraction error: {e}")
        return ""


@router.post("/parse")
async def parse_bidding_document(file: UploadFile = File(...)):
    """上传招标文件 → AI 解析提取关键信息"""
    
    try:
        # 1. 读取文件内容
        file_bytes = await file.read()
        filename = file.filename or "unknown"
        logger.info(f"Parsing bidding document: {filename} ({len(file_bytes)} bytes)")
        
        # 2. 提取文本
        if filename.endswith('.docx'):
            text = extract_text_from_docx(file_bytes)
        elif filename.endswith('.txt'):
            text = file_bytes.decode('utf-8', errors='ignore')
        else:
            return {"success": False, "message": f"不支持的文件格式: {filename}，请上传 .docx 或 .txt 文件"}
        
        if not text or len(text) < 50:
            return {"success": False, "message": "文件内容提取失败或内容过少"}
        
        # 截取前 8000 字符避免超出上下文限制
        text_truncated = text[:8000]
        logger.info(f"Extracted {len(text)} chars, using first {len(text_truncated)} for parsing")
        
        # 3. 调用 Qwen 解析
        llm = get_llm()
        prompt = f"请分析以下招标文件内容，提取关键信息：\n\n{text_truncated}"
        
        response = await llm.generate(prompt, system=PARSE_SYSTEM_PROMPT)
        
        # 4. 尝试解析 JSON
        try:
            # 提取 JSON 块（处理可能包含 markdown 代码块的情况）
            json_match = re.search(r'\{[\s\S]*\}', response)
            if json_match:
                parsed = json.loads(json_match.group())
            else:
                parsed = json.loads(response)
        except json.JSONDecodeError:
            logger.warning(f"JSON parse failed, returning raw: {response[:200]}")
            parsed = {"raw_response": response, "parse_error": True}
        
        return {
            "success": True,
            "data": {
                "parsed": parsed,
                "text_length": len(text),
                "filename": filename,
                "model": llm.get_model_name(),
            }
        }
        
    except Exception as e:
        logger.error(f"Parse error: {e}")
        return {"success": False, "message": str(e)}


# ─────────────────────────────────────────────
#  2. 生成端点：基于解析结果 + 律师信息生成标书
# ─────────────────────────────────────────────

BIDDING_SYSTEM_PROMPT = """你是一位资深的律所投标文书专家，专门为律师事务所生成投标文件框架。
你的任务是根据用户提供的投标人信息和项目背景，按照中国移动等大型央企的招标文件格式要求，生成完整的投标文件框架（商务分册部分）。

## 格式规范（严格执行）

你输出的内容会被 Markdown 渲染器渲染成 Word 风格文档，因此必须严格遵守以下格式要求：

1. **标题层级**：
   - `#` = 文档总标题（仅用一次，如 `# 投标文件（商务分册）`）
   - `##` = 章节标题（如 `## 一、投标函`、`## 二、资格审查资料`）
   - `###` = 小节标题（如 `### 2.1 投标人基本情况表`）
2. **表格**：资格审查表、投标一览表等必须使用 Markdown 表格 `| col | col |` 格式
3. **占位符**：所有需要填写的数据用 **【待填写】** 标注，附件位置用 **【此处附XX扫描件】** 标注
4. **章节分隔**：每个大章节（##）之间用 `---` 分隔线
5. **正文段落**：直接写文字，不要在普通段落前加 `-` 或 `*` 符号
6. **金额**：统一为人民币元
7. **日期**：XXXX年XX月XX日
8. **签章区域**：每章结尾如有签章需求，用表格或加粗标注签署区

## 章节结构（按顺序生成）

1. **投标函** — 致招标人的正式承诺函，包含投标报价、有效期、联系方式等
2. **资格审查资料** — 投标人基本情况表（Markdown 表格）
3. **投标一览表** — 服务内容/报价明细表（Markdown 表格）
4. **企业信誉声明函** — 无行贿/违法记录声明
5. **非联合体投标及不转包承诺函** — 独立投标声明
6. **商务条款偏离表** — 逐条列出"无偏离"（Markdown 表格）
7. **法定代表人身份证明** — 含姓名/性别/年龄/职务
8. **法定代表人授权委托书** — 如有委托代理人，含被授权人信息和权限范围

## 注意事项
- 已知信息直接填入，未知信息用 **【待填写】** 标注
- 如果提供了招标文件解析信息（关键要求/否决条件），要在文档中针对性响应
- 内容要充实饱满，每章不少于 3-5 个段落
- 投标函中的承诺条款要完整、规范"""


class BiddingRequest(BaseModel):
    """投标文件生成请求"""
    company_name: str
    legal_representative: str
    project_name: str
    client_name: str
    project_id: Optional[str] = ""
    registered_capital: Optional[str] = ""
    established_date: Optional[str] = ""
    address: Optional[str] = ""
    contact_person: Optional[str] = ""
    contact_phone: Optional[str] = ""
    contact_email: Optional[str] = ""
    bid_amount: Optional[str] = ""
    guarantee_amount: Optional[str] = ""
    delegate_name: Optional[str] = ""
    validity_days: Optional[str] = "120"
    # 来自招标文件解析的额外上下文
    parsed_requirements: Optional[str] = ""
    parsed_risks: Optional[str] = ""
    budget: Optional[str] = ""
    stream: bool = True


@router.post("/generate")
async def generate_bidding_document(req: BiddingRequest):
    """生成投标文件框架 — SSE 流式输出"""

    # 构建 prompt
    user_prompt = f"""请根据以下信息，生成完整的投标文件框架（商务分册）：

## 投标基本信息

- **投标人名称**：{req.company_name}
- **法定代表人**：{req.legal_representative}
- **项目名称**：{req.project_name}
- **招标人名称**：{req.client_name}
- **项目编号**：{req.project_id or '【待填写】'}
- **注册资本**：{req.registered_capital or '【待填写】'}
- **成立时间**：{req.established_date or '【待填写】'}
- **注册地址**：{req.address or '【待填写】'}
- **联系人**：{req.contact_person or '【待填写】'}
- **联系电话**：{req.contact_phone or '【待填写】'}
- **电子邮箱**：{req.contact_email or '【待填写】'}
- **投标报价**：{req.bid_amount or '【待填写】'}
- **投标保证金**：{req.guarantee_amount or '【待填写】'}
- **委托代理人**：{req.delegate_name or '无（法定代表人直接投标）'}
- **投标有效期**：{req.validity_days}天
- **项目预算/最高限价**：{req.budget or '【待填写】'}"""

    # 如果有从招标文件中解析出的要求，加入 prompt
    if req.parsed_requirements:
        user_prompt += f"\n\n## 招标文件关键要求\n\n{req.parsed_requirements}"
    
    if req.parsed_risks:
        user_prompt += f"\n\n## 否决条件（务必在投标文件中避免）\n\n{req.parsed_risks}"

    user_prompt += "\n\n请按照标准格式生成完整的投标文件框架，已知信息直接填入，未知信息用【待填写】标注。"

    try:
        llm = get_llm()
        logger.info(f"Bidding generation: {req.company_name} → {req.project_name}")

        if req.stream:
            async def event_generator():
                try:
                    async for chunk in llm.stream(user_prompt, system=BIDDING_SYSTEM_PROMPT):
                        yield f"data: {json.dumps({'content': chunk, 'done': False}, ensure_ascii=False)}\n\n"
                    yield f"data: {json.dumps({'content': '', 'done': True}, ensure_ascii=False)}\n\n"
                except Exception as e:
                    logger.error(f"Bidding stream error: {e}")
                    yield f"data: {json.dumps({'content': f'[生成错误] {str(e)}', 'done': True}, ensure_ascii=False)}\n\n"

            return StreamingResponse(event_generator(), media_type="text/event-stream")
        else:
            response = await llm.generate(user_prompt, system=BIDDING_SYSTEM_PROMPT)
            return {"success": True, "data": {"content": response, "model": llm.get_model_name()}}

    except Exception as e:
        logger.error(f"Bidding generation error: {e}")
        return {"success": False, "message": str(e)}


# ─────────────────────────────────────────────
#  3. 新版完整投标管线 (Phase 1)
# ─────────────────────────────────────────────

import os
import time
import asyncio
from typing import Dict, Any, List

from app.core.skills.builtin.tender_parsing import TenderParsingSkill
from app.core.skills.builtin.requirement_extraction import RequirementExtractionSkill
from app.core.skills.builtin.content_generation import ContentGenerationSkill
from app.core.skills.builtin.template_filling import TemplateFillingSkill
from app.core.skills.builtin.docx_assembly import DocxAssemblySkill
from app.core.skills.builtin.rule_verification import RuleVerificationSkill
from app.core.skills.builtin.template_store import TemplateStoreSkill
from app.core.skills.builtin.data_retrieval import DataRetrievalSkill
from app.core.rag.tender_index import TenderIndex
from app.config import settings

# In-memory storage for bidding tasks (production would use DB)
_bidding_tasks: Dict[str, Dict[str, Any]] = {}

# Skill instances
_parser = TenderParsingSkill()
_extractor = RequirementExtractionSkill()
_generator = ContentGenerationSkill()
_filler = TemplateFillingSkill()
_assembler = DocxAssemblySkill()
_verifier = RuleVerificationSkill()
_template_store = TemplateStoreSkill()
_data_retrieval = DataRetrievalSkill()


class FullBiddingRequest(BaseModel):
    """Request for full bidding pipeline."""
    company_data: Optional[Dict[str, str]] = None
    llm_provider: str = "qwen"


@router.post("/parse-structure")
async def parse_tender_structure(file: UploadFile = File(...),
                                  llm_provider: str = Form("qwen")):
    """上传招标文件 → SSE 流式解析 → 实时输出进度

    Returns SSE stream with progress events, final event contains the full result.
    """
    file_bytes = await file.read()
    filename = file.filename or "unknown.docx"
    logger.info(f"[Full Pipeline] Parse structure (stream): {filename} ({len(file_bytes)} bytes), LLM={llm_provider}")

    if not filename.endswith('.docx'):
        async def _err():
            yield f"data: {json.dumps({'type': 'error', 'message': '目前仅支持 .docx 格式招标文件'})}\n\n"
        return StreamingResponse(_err(), media_type="text/event-stream")

    async def _stream_parse():
        import asyncio

        def emit(event_type, **kwargs):
            payload = {"type": event_type, **kwargs}
            return f"data: {json.dumps(payload, ensure_ascii=False)}\n\n"

        try:
            # ── Step 0: Save file ──
            yield emit("log", message=f"📄 收到文件: {filename} ({len(file_bytes)/1024:.0f}KB)")
            os.makedirs(settings.UPLOAD_DIR, exist_ok=True)
            temp_path = os.path.join(settings.UPLOAD_DIR, f"tender_{int(time.time())}_{filename}")
            with open(temp_path, "wb") as f:
                f.write(file_bytes)
            yield emit("log", message=f"💾 文件已保存")

            # ── Step 1: Parse document structure ──
            yield emit("phase", phase="parsing", message="🔍 正在解析文档结构 (python-docx)...")
            parse_result = await _parser.execute({"file_path": temp_path})
            total_sections = parse_result.get("total_sections", 0)
            text_len = len(parse_result.get("raw_text", ""))
            yield emit("log", message=f"✅ 文档解析完成: {total_sections} 个段落, {text_len} 字符")

            # Show first few section titles
            raw_sections = parse_result.get("sections", [])
            sample_titles = [s.get("title", "") for s in raw_sections[:8] if s.get("title")]
            if sample_titles:
                yield emit("log", message="📋 发现的章节标题 (前8个):")
                for i, t in enumerate(sample_titles, 1):
                    yield emit("log", message=f"   {i}. {t[:60]}")

            # ── Step 2: LLM Classification ──
            yield emit("phase", phase="classifying", message=f"🤖 正在用 AI 分类 {total_sections} 个章节...")
            extract_result = await _extractor.execute({
                "raw_text": parse_result["raw_text"],
                "sections": parse_result["sections"],
                "llm_provider": llm_provider,
            })

            # Show classified structure
            volumes = extract_result.get("volumes", [])
            total_secs = sum(len(v.get("sections", [])) for v in volumes)
            yield emit("log", message=f"✅ AI 分类完成: {len(volumes)} 个分册, {total_secs} 个有效章节")

            for vol in volumes:
                vol_name = vol.get("name", "?")
                vol_secs = vol.get("sections", [])
                yield emit("log", message=f"📁 {vol_name} ({len(vol_secs)} 章节)")
                for sec in vol_secs[:5]:
                    sec_type = sec.get("type", "?")
                    icon = {"narrative": "📝", "table": "📊", "form": "📋", "qualification": "🏅"}.get(sec_type, "📄")
                    yield emit("section", title=sec.get("title", "?"), type=sec_type, icon=icon)
                if len(vol_secs) > 5:
                    yield emit("log", message=f"   ... 还有 {len(vol_secs) - 5} 个章节")

            # ── Step 3: Build vector index ──
            yield emit("phase", phase="indexing", message="🔗 正在构建检索索引 (BGE)...")
            tender_index = TenderIndex()
            try:
                chunk_count = tender_index.build(
                    raw_text=parse_result["raw_text"],
                    sections=parse_result["sections"],
                )
                yield emit("log", message=f"✅ 检索索引构建完成: {chunk_count} 个文本块")
            except Exception as e:
                logger.warning(f"Tender index build failed: {e}")
                tender_index = None
                yield emit("log", message=f"⚠️ 检索索引构建失败 (可继续生成): {str(e)[:50]}")

            # ── Store task ──
            task_id = f"bid_{int(time.time())}"

            # Structure verification
            structure_warnings = []
            if tender_index and tender_index.is_built:
                try:
                    all_titles = [sec.get("title", "")
                                  for vol in volumes
                                  for sec in vol.get("sections", [])]
                    verification = tender_index.verify_structure(all_titles)
                    uncovered = [v for v in verification if not v["covered"]]
                    if uncovered:
                        structure_warnings = [
                            f"招标要求 '{v['requirement']}' 可能未覆盖"
                            for v in uncovered[:5]
                        ]
                        yield emit("log", message=f"⚠️ 发现 {len(uncovered)} 项招标要求可能未覆盖")
                except Exception:
                    pass

            _bidding_tasks[task_id] = {
                "task_id": task_id,
                "status": "parsed",
                "tender_file": temp_path,
                "parse_result": parse_result,
                "requirements": extract_result,
                "tender_index": tender_index,
                "generated_sections": [],
                "verification": None,
                "output_file": None,
                "created_at": time.time(),
            }

            yield emit("log", message=f"🎉 解析全部完成! 任务ID: {task_id}")

            # ── Final event: complete with full data ──
            yield emit("complete",
                        task_id=task_id,
                        requirements=extract_result,
                        raw_sections_count=parse_result["total_sections"],
                        text_length=text_len,
                        index_chunks=tender_index.chunk_count if tender_index else 0,
                        structure_warnings=structure_warnings)

        except Exception as e:
            logger.error(f"[Full Pipeline] Parse stream error: {e}", exc_info=True)
            yield emit("error", message=str(e))

    return StreamingResponse(_stream_parse(), media_type="text/event-stream")


@router.post("/generate-full/{task_id}")
async def generate_full_document(task_id: str, req: FullBiddingRequest):
    """逐章节生成完整投标文件 — SSE 流式进度

    Reads the requirements from parse_structure step, generates each section,
    assembles into .docx, and runs verification.
    """
    task = _bidding_tasks.get(task_id)
    if not task:
        return {"success": False, "message": f"任务 {task_id} 不存在"}

    requirements = task.get("requirements", {})
    company_data = req.company_data or {}
    llm_provider = req.llm_provider

    # Get company info summary for prompts
    company_info = TemplateFillingSkill.get_company_info_summary(company_data)

    async def event_generator():
        generated_sections = []
        total_sections = 0
        current = 0

        # Collect all sections from all volumes
        all_sections = []
        for volume in requirements.get("volumes", []):
            for section in volume.get("sections", []):
                all_sections.append(section)
        total_sections = len(all_sections)

        if total_sections == 0:
            yield _sse({"type": "error", "message": "未找到需要生成的章节"})
            return

        # ── Step 3: Template matching ──
        matched_skeletons = {}
        match_result = _template_store.match_template({"requirements": requirements})
        if match_result.get("matched"):
            matched_skeletons = match_result.get("skeletons", {})
            yield _sse({
                "type": "template_matched",
                "template_name": match_result.get("template_name", ""),
                "score": match_result.get("score", 0),
                "skeleton_count": len(matched_skeletons),
            })

        # Count how many will use LLM vs code templates
        llm_count = sum(1 for s in all_sections if s.get("type") == "narrative")
        code_count = total_sections - llm_count

        yield _sse({
            "type": "start",
            "total_sections": total_sections,
            "llm_sections": llm_count,
            "code_sections": code_count,
            "message": f"开始生成投标文件：{code_count} 个章节用代码模板，{llm_count} 个用LLM",
        })

        # ── Concurrent generation (5 workers) ──
        import asyncio

        CONCURRENCY = 5
        sem = asyncio.Semaphore(CONCURRENCY)
        event_queue: asyncio.Queue = asyncio.Queue()
        completed_count = 0

        async def _gen_one(idx: int, section: dict):
            """Generate a single section with semaphore limiting."""
            title = section.get("title", f"章节{idx + 1}")
            sec_type = section.get("type", "narrative")

            await event_queue.put(_sse({
                "type": "progress",
                "current": idx + 1,
                "total": total_sections,
                "section_title": title,
                "section_type": sec_type,
                "status": "generating",
                "method": "llm" if sec_type == "narrative" else "template",
            }))

            async with sem:
                try:
                    skeleton_info = matched_skeletons.get(title, {})
                    skeleton_text = skeleton_info.get("skeleton") if skeleton_info else None

                    reference_data = "暂无参考资料"
                    tender_index = task.get("tender_index")
                    if sec_type == "narrative" and tender_index and tender_index.is_built:
                        query = f"{title} {section.get('content_hints', '')}"
                        relevant_chunks = tender_index.search(query, top_k=5)
                        if relevant_chunks:
                            reference_data = "\n\n---\n\n".join(relevant_chunks)

                    result = await _generator.execute({
                        "section": section,
                        "company_info": company_info,
                        "reference_data": reference_data,
                        "llm_provider": llm_provider,
                        "skeleton": skeleton_text,
                    })

                    result["order"] = section.get("order", idx + 1)
                    result["type"] = section.get("type", "narrative")
                    result["level"] = 2

                    await event_queue.put(_sse({
                        "type": "section_done",
                        "current": idx + 1,
                        "total": total_sections,
                        "section_title": title,
                        "status": result.get("status", "generated"),
                        "content_length": len(result.get("content", "")),
                        "missing_fields": result.get("missing_fields", []),
                    }))
                    return result

                except Exception as e:
                    logger.error(f"Error generating section '{title}': {e}")
                    await event_queue.put(_sse({
                        "type": "section_error",
                        "current": idx + 1,
                        "section_title": title,
                        "error": str(e),
                    }))
                    return {
                        "title": title,
                        "content": f"[生成失败：{str(e)}]",
                        "order": section.get("order", idx + 1),
                        "type": section.get("type", "narrative"),
                        "level": 2,
                        "missing_fields": [],
                        "status": "error",
                    }

        # Launch all tasks, drain SSE queue while waiting
        tasks = [asyncio.create_task(_gen_one(i, sec)) for i, sec in enumerate(all_sections)]

        done_count = 0
        while done_count < len(tasks):
            # Check for newly queued events
            try:
                event = await asyncio.wait_for(event_queue.get(), timeout=0.2)
                yield event
            except asyncio.TimeoutError:
                pass
            # Count finished tasks
            done_count = sum(1 for t in tasks if t.done())

        # Drain remaining events in queue
        while not event_queue.empty():
            yield await event_queue.get()

        # Collect results in original order
        generated_sections = [t.result() for t in tasks]

        # Assemble document
        yield _sse({"type": "assembling", "message": "正在组装 Word 文档..."})

        try:
            assembly_result = await _assembler.execute({
                "bid_title": requirements.get("bid_title", "投标文件"),
                "sections": generated_sections,
                "format_rules": requirements.get("format_requirements", {}),
            })

            task["output_file"] = assembly_result["file_path"]
            task["output_filename"] = assembly_result["filename"]
        except Exception as e:
            logger.error(f"Assembly error: {e}")
            assembly_result = {"file_path": None, "error": str(e)}

        # Run verification
        yield _sse({"type": "verifying", "message": "正在校验投标文件..."})

        try:
            verification = await _verifier.execute({
                "tender_requirements": requirements,
                "generated_sections": generated_sections,
            })
            task["verification"] = verification
        except Exception as e:
            logger.error(f"Verification error: {e}")
            verification = {"overall_status": "ERROR", "error": str(e)}

        # Update task
        task["status"] = "completed"
        task["generated_sections"] = generated_sections

        yield _sse({
            "type": "complete",
            "task_id": task_id,
            "file_path": assembly_result.get("file_path"),
            "filename": assembly_result.get("filename"),
            "page_estimate": assembly_result.get("page_count_estimate", 0),
            "section_count": len(generated_sections),
            "verification": verification,
        })

    return StreamingResponse(event_generator(), media_type="text/event-stream")


@router.post("/verify/{task_id}")
async def verify_document(task_id: str):
    """对已生成的投标文件重新运行校验"""
    task = _bidding_tasks.get(task_id)
    if not task:
        return {"success": False, "message": f"任务 {task_id} 不存在"}

    if not task.get("generated_sections"):
        return {"success": False, "message": "尚未生成投标文件"}

    try:
        verification = await _verifier.execute({
            "tender_requirements": task.get("requirements", {}),
            "generated_sections": task["generated_sections"],
        })
        task["verification"] = verification

        return {"success": True, "data": verification}
    except Exception as e:
        return {"success": False, "message": str(e)}


@router.get("/download/{task_id}")
async def download_document(task_id: str):
    """下载生成的投标文件 .docx"""
    from fastapi.responses import FileResponse

    task = _bidding_tasks.get(task_id)
    if not task:
        return {"success": False, "message": f"任务 {task_id} 不存在"}

    file_path = task.get("output_file")
    if not file_path or not os.path.exists(file_path):
        return {"success": False, "message": "文件尚未生成或已被删除"}

    filename = task.get("output_filename", "bid_document.docx")
    return FileResponse(
        path=file_path,
        filename=filename,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


@router.get("/tasks")
async def list_tasks():
    """列出所有投标任务"""
    tasks = []
    for tid, t in _bidding_tasks.items():
        tasks.append({
            "task_id": tid,
            "status": t.get("status"),
            "created_at": t.get("created_at"),
            "section_count": len(t.get("generated_sections", [])),
            "has_output": t.get("output_file") is not None,
        })
    return {"success": True, "data": tasks}


# ── Template management endpoints ──

@router.get("/templates")
async def list_templates():
    """列出所有投标模板"""
    result = _template_store.list_templates()
    return {"success": True, "data": result}


@router.get("/templates/{template_id}")
async def get_template(template_id: str):
    """获取单个模板详情"""
    result = _template_store.get_template(template_id)
    if "error" in result:
        return {"success": False, "message": result["error"]}
    return {"success": True, "data": result}


@router.post("/templates/save/{task_id}")
async def save_as_template(task_id: str):
    """将已完成的投标任务存为模板"""
    task = _bidding_tasks.get(task_id)
    if not task:
        return {"success": False, "message": f"任务 {task_id} 不存在"}
    if not task.get("generated_sections"):
        return {"success": False, "message": "尚未生成投标文件"}

    result = _template_store.save_template({
        "requirements": task.get("requirements", {}),
        "generated_sections": task["generated_sections"],
        "source_file": task.get("tender_file", "unknown"),
    })
    return {"success": True, "data": result}


@router.delete("/templates/{template_id}")
async def delete_template(template_id: str):
    """删除模板"""
    result = _template_store.delete_template(template_id)
    if "error" in result:
        return {"success": False, "message": result["error"]}
    return {"success": True, "data": result}


# ── Historical bid document & materials endpoints (S2/S3) ──

@router.post("/upload-historical")
async def upload_historical_bid(
    file: UploadFile = File(...),
    llm_provider: str = Form("qwen"),
):
    """上传历史投标/素材文件 → 提取 → 返回变更对比 (不自动入库)

    Step 1 of 2-step flow:
    1. upload-historical → extract + diff (this endpoint)
    2. confirm-materials → save to store (user confirms)
    """
    import time as _time

    # Save uploaded file
    filename = f"historical_{int(_time.time())}_{file.filename}"
    file_path = os.path.join("uploads", filename)
    os.makedirs("uploads", exist_ok=True)
    with open(file_path, "wb") as f:
        content = await file.read()
        f.write(content)

    logger.info(f"Historical bid uploaded: {file_path} ({len(content)} bytes)")

    try:
        # Parse and extract materials
        from app.core.skills.builtin.bid_document_parser import BidDocumentParserSkill
        parser = BidDocumentParserSkill()
        materials = await parser.execute({
            "file_path": file_path,
            "llm_provider": llm_provider,
        })

        # Inject source file info into every extracted item
        original_filename = file.filename  # user's original filename
        for category in ("resumes", "projects", "qualifications", "narrative_chunks"):
            for item in materials.get(category, []):
                item["_source_file"] = original_filename
                item["_source_path"] = file_path

        # Diff against existing store (don't save yet)
        from app.core.skills.builtin.material_store import get_material_store
        store = get_material_store()
        diff = store.diff_materials(materials)

        # Store pending materials in memory for confirmation
        upload_id = f"upload_{int(_time.time())}"
        _pending_uploads[upload_id] = materials

        return {
            "success": True,
            "data": {
                "upload_id": upload_id,
                "source_file": filename,
                "extracted": {
                    "resumes": len(materials.get("resumes", [])),
                    "projects": len(materials.get("projects", [])),
                    "qualifications": len(materials.get("qualifications", [])),
                    "narrative_chunks": len(materials.get("narrative_chunks", [])),
                },
                "diff": diff,
                "materials": {
                    "resumes": materials.get("resumes", []),
                    "projects": materials.get("projects", []),
                    "qualifications": materials.get("qualifications", []),
                },
            }
        }

    except Exception as e:
        logger.error(f"Historical bid processing failed: {e}")
        return {"success": False, "message": f"处理失败: {str(e)}"}


# In-memory pending uploads (awaiting user confirmation)
_pending_uploads = {}  # type: Dict[str, Any]


class ConfirmMaterialsRequest(BaseModel):
    upload_id: str
    # Optional: subset of items to save (if user deselects some)
    # If empty, save all extracted materials
    selected_resumes: Optional[list] = None
    selected_projects: Optional[list] = None
    selected_qualifications: Optional[list] = None


@router.post("/confirm-materials")
async def confirm_materials(req: ConfirmMaterialsRequest):
    """用户确认后，将提取的素材入库

    Step 2 of 2-step flow. Called after user reviews diff from upload-historical.
    """
    materials = _pending_uploads.pop(req.upload_id, None)
    if not materials:
        return {"success": False, "message": f"上传 {req.upload_id} 不存在或已过期"}

    # If user selected specific items, filter
    if req.selected_resumes is not None:
        names = set(req.selected_resumes)
        materials["resumes"] = [r for r in materials.get("resumes", [])
                                if r.get("name") in names]
    if req.selected_projects is not None:
        names = set(req.selected_projects)
        materials["projects"] = [p for p in materials.get("projects", [])
                                 if p.get("project_name") in names]
    if req.selected_qualifications is not None:
        names = set(req.selected_qualifications)
        materials["qualifications"] = [q for q in materials.get("qualifications", [])
                                       if q.get("name") in names]

    from app.core.skills.builtin.material_store import get_material_store
    store = get_material_store()
    save_counts = store.save_materials(materials)

    return {
        "success": True,
        "data": {
            "saved": save_counts,
            "total_store": store.get_summary(),
        }
    }


@router.get("/materials")
async def get_materials():
    """获取所有已提取的素材"""
    from app.core.skills.builtin.material_store import get_material_store
    store = get_material_store()
    return {"success": True, "data": store.get_all_materials()}


@router.get("/materials/summary")
async def get_materials_summary():
    """获取素材库概要统计"""
    from app.core.skills.builtin.material_store import get_material_store
    store = get_material_store()
    return {"success": True, "data": store.get_summary()}


@router.get("/materials/search")
async def search_materials(
    q: str = "",
    type: str = "all",
):
    """搜索素材库

    Query params:
        q: 搜索关键词
        type: resumes | projects | qualifications | narratives | all
    """
    from app.core.skills.builtin.material_store import get_material_store
    store = get_material_store()

    results = {}
    if type in ("all", "resumes"):
        results["resumes"] = store.search_resumes(query=q)
    if type in ("all", "projects"):
        results["projects"] = store.search_projects(query=q)
    if type in ("all", "qualifications"):
        results["qualifications"] = store.get_qualifications()
    if type in ("all", "narratives") and q:
        results["narratives"] = await store.search_narratives(query=q)

    return {"success": True, "data": results}


# ── Material CRUD Endpoints ──

class MaterialUpdateRequest(BaseModel):
    updates: dict


@router.put("/materials/resumes/{name}")
async def update_resume(name: str, req: MaterialUpdateRequest):
    """更新律师简历"""
    from app.core.skills.builtin.material_store import get_material_store
    store = get_material_store()
    ok = store.update_resume(name, req.updates)
    if ok:
        return {"success": True}
    return {"success": False, "message": f"简历 '{name}' 未找到"}


@router.delete("/materials/resumes/{name}")
async def delete_resume(name: str):
    """删除律师简历"""
    from app.core.skills.builtin.material_store import get_material_store
    store = get_material_store()
    ok = store.delete_resume(name)
    if ok:
        return {"success": True}
    return {"success": False, "message": f"简历 '{name}' 未找到"}


@router.put("/materials/projects/{project_name}")
async def update_project(project_name: str, req: MaterialUpdateRequest):
    """更新项目业绩"""
    from app.core.skills.builtin.material_store import get_material_store
    store = get_material_store()
    ok = store.update_project(project_name, req.updates)
    if ok:
        return {"success": True}
    return {"success": False, "message": f"项目 '{project_name}' 未找到"}


@router.delete("/materials/projects/{project_name}")
async def delete_project(project_name: str):
    """删除项目业绩"""
    from app.core.skills.builtin.material_store import get_material_store
    store = get_material_store()
    projects = store.get_projects()
    new_projects = [p for p in projects if p.get("project_name") != project_name]
    if len(new_projects) < len(projects):
        store._save_json(store.projects_file, new_projects)
        return {"success": True}
    return {"success": False, "message": f"项目 '{project_name}' 未找到"}


@router.delete("/materials/qualifications/{name}")
async def delete_qualification(name: str):
    """删除资质证书"""
    from app.core.skills.builtin.material_store import get_material_store
    store = get_material_store()
    quals = store.get_qualifications()
    new_quals = [q for q in quals if q.get("name") != name]
    if len(new_quals) < len(quals):
        store._save_json(store.qualifications_file, new_quals)
        return {"success": True}
    return {"success": False, "message": f"资质 '{name}' 未找到"}


class AddMaterialRequest(BaseModel):
    type: str  # resumes | projects | qualifications
    data: dict


@router.post("/materials/add")
async def add_material(req: AddMaterialRequest):
    """手动新增素材"""
    from app.core.skills.builtin.material_store import get_material_store
    store = get_material_store()

    if req.type == "resumes":
        items = store.get_resumes()
        items.append(req.data)
        items = store._dedup_by_field(items, "name")
        store._save_json(store.resumes_file, items)
    elif req.type == "projects":
        items = store.get_projects()
        items.append(req.data)
        items = store._dedup_by_field(items, "project_name")
        store._save_json(store.projects_file, items)
    elif req.type == "qualifications":
        items = store.get_qualifications()
        items.append(req.data)
        items = store._dedup_by_field(items, "name")
        store._save_json(store.qualifications_file, items)
    else:
        return {"success": False, "message": f"未知类型: {req.type}"}

    return {"success": True, "data": store.get_summary()}


@router.get("/materials/source-files/{name}")
async def get_source_files(name: str):
    """查找某个人/项目/资质对应的原始素材文件

    1. 先检查素材数据中的 _source_file/_source_path (精确来源)
    2. 再扫描 client_materials + uploads 目录做文件名模糊匹配
    """
    from app.core.skills.builtin.material_store import get_material_store
    store = get_material_store()

    matches = []
    seen_paths = set()

    # Priority 1: Check _source_path in material data
    for category_getter in [store.get_resumes, store.get_projects, store.get_qualifications]:
        for item in category_getter():
            item_name = item.get("name", "") or item.get("project_name", "") or ""
            if not item_name:
                continue
            if name in item_name or item_name in name:
                src_path = item.get("_source_path", "")
                src_file = item.get("_source_file", "")
                if src_path and os.path.exists(src_path) and src_path not in seen_paths:
                    seen_paths.add(src_path)
                    size = os.path.getsize(src_path)
                    ext = os.path.splitext(src_file)[1].lower()
                    matches.append({
                        "filename": src_file,
                        "folder": "上传素材",
                        "relative_path": src_path,
                        "size_bytes": size,
                        "size_display": f"{size / 1024 / 1024:.1f}MB" if size > 1024 * 1024 else f"{size / 1024:.0f}KB",
                        "file_type": ext.lstrip("."),
                        "source": "uploaded",
                    })

    # Priority 2: Scan client_materials ONLY for original source file
    client_dir = os.path.join("uploads", "client_materials")
    if os.path.isdir(client_dir):
        for root, dirs, files in os.walk(client_dir):
            for fname in files:
                if fname.startswith("."):
                    continue
                if name in fname:
                    full_path = os.path.join(root, fname)
                    if full_path in seen_paths:
                        continue
                    seen_paths.add(full_path)
                    rel_path = os.path.relpath(full_path)
                    folder = os.path.basename(root)
                    try:
                        size = os.path.getsize(full_path)
                    except OSError:
                        size = 0
                    ext = os.path.splitext(fname)[1].lower()
                    matches.append({
                        "filename": fname,
                        "folder": folder,
                        "relative_path": rel_path,
                        "size_bytes": size,
                        "size_display": f"{size / 1024 / 1024:.1f}MB" if size > 1024 * 1024 else f"{size / 1024:.0f}KB",
                        "file_type": ext.lstrip("."),
                    })

    return {"success": True, "data": {"name": name, "files": matches}}


@router.post("/materials/backfill-sources")
async def backfill_sources():
    """回填已有素材的 _source_file 字段

    扫描 client_materials 目录，根据素材名称匹配文件名，
    自动为没有 _source_file 的素材补充来源信息。
    """
    client_dir = os.path.join("uploads", "client_materials")
    if not os.path.isdir(client_dir):
        return {"success": False, "message": "client_materials 目录不存在"}

    # Collect all files
    all_files = []
    for root, dirs, files in os.walk(client_dir):
        for fname in files:
            if fname.startswith("."):
                continue
            full_path = os.path.join(root, fname)
            folder = os.path.basename(root)
            all_files.append({
                "filename": fname,
                "folder": folder,
                "full_path": full_path,
            })

    from app.core.skills.builtin.material_store import get_material_store
    store = get_material_store()
    updated = {"resumes": 0, "projects": 0, "qualifications": 0}

    # Backfill resumes
    resumes = store.get_resumes()
    for r in resumes:
        if r.get("_source_file"):
            continue
        name = r.get("name", "")
        if not name:
            continue
        matching = [f for f in all_files if name in f["filename"]]
        if matching:
            r["_source_file"] = matching[0]["filename"]
            r["_source_path"] = matching[0]["full_path"]
            r["_source_files"] = [f["filename"] for f in matching]
            updated["resumes"] += 1
    store._save_json(store.resumes_file, resumes)

    # Backfill projects
    projects = store.get_projects()
    for p in projects:
        if p.get("_source_file"):
            continue
        pname = p.get("project_name", "")
        client = p.get("client", "")
        matching = [f for f in all_files
                    if (pname and pname in f["filename"]) or (client and client in f["filename"])]
        if matching:
            p["_source_file"] = matching[0]["filename"]
            p["_source_path"] = matching[0]["full_path"]
            updated["projects"] += 1
    store._save_json(store.projects_file, projects)

    # Backfill qualifications
    quals = store.get_qualifications()
    for q in quals:
        if q.get("_source_file"):
            continue
        qname = q.get("name", "")
        if not qname:
            continue
        matching = [f for f in all_files if qname in f["filename"]]
        if matching:
            q["_source_file"] = matching[0]["filename"]
            q["_source_path"] = matching[0]["full_path"]
            updated["qualifications"] += 1
    store._save_json(store.qualifications_file, quals)

    return {"success": True, "data": {"updated": updated}}


@router.get("/materials/preview")
async def preview_material_file(path: str):
    """在线预览素材文件

    支持 docx (转 HTML), pdf, 图片 (直接返回)。
    path 参数是相对于 uploads/client_materials 的路径。
    """
    import os

    # Handle two path formats:
    # 1. Old: relative to client_materials (e.g. "kindofpdfword/团队人员资料/file.docx")
    # 2. New: relative to CWD (e.g. "uploads/historical_123_file.docx")
    if path.startswith("uploads/") or path.startswith("uploads\\"):
        full_path = os.path.normpath(path)
    else:
        client_dir = os.path.join("uploads", "client_materials")
        full_path = os.path.normpath(os.path.join(client_dir, path))

    # Security: prevent path traversal (must stay under uploads/)
    if not full_path.startswith("uploads"):
        return HTMLResponse("<h1>403 Forbidden</h1>", status_code=403)

    if not os.path.isfile(full_path):
        return HTMLResponse(f"<h1>404 文件不存在</h1><p>{path}</p>", status_code=404)

    ext = os.path.splitext(full_path)[1].lower()
    filename = os.path.basename(full_path)

    # ── DOCX → styled HTML ──
    if ext in (".docx", ".doc"):
        try:
            from docx import Document
            doc = Document(full_path)

            # Extract images from docx to base64
            import base64
            image_map = {}
            for rel in doc.part.rels.values():
                if "image" in rel.reltype:
                    img_data = rel.target_part.blob
                    content_type = rel.target_part.content_type or "image/png"
                    b64 = base64.b64encode(img_data).decode()
                    image_map[rel.rId] = f"data:{content_type};base64,{b64}"

            # Build HTML
            html_parts = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    html_parts.append("<br/>")
                    continue
                style = para.style.name if para.style else ""
                if "Heading 1" in style:
                    html_parts.append(f"<h1>{text}</h1>")
                elif "Heading 2" in style:
                    html_parts.append(f"<h2>{text}</h2>")
                elif "Heading 3" in style:
                    html_parts.append(f"<h3>{text}</h3>")
                else:
                    # Check for bold/emphasis
                    if para.runs and all(r.bold for r in para.runs if r.text.strip()):
                        html_parts.append(f"<p><strong>{text}</strong></p>")
                    else:
                        html_parts.append(f"<p>{text}</p>")

            # Extract tables
            for table in doc.tables:
                html_parts.append('<table>')
                for ri, row in enumerate(table.rows):
                    html_parts.append('<tr>')
                    tag = 'th' if ri == 0 else 'td'
                    for cell in row.cells:
                        html_parts.append(f'<{tag}>{cell.text}</{tag}>')
                    html_parts.append('</tr>')
                html_parts.append('</table>')

            # Render images
            for rid, data_url in image_map.items():
                html_parts.append(f'<img src="{data_url}" style="max-width:100%;margin:12px 0;border-radius:4px" />')

            body = "\n".join(html_parts)

            page = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8"/>
<title>{filename}</title>
<style>
  * {{ margin: 0; padding: 0; box-sizing: border-box; }}
  body {{
    font-family: -apple-system, "PingFang SC", "Microsoft YaHei", sans-serif;
    background: #1a1a2e; color: #e0e0e0;
    padding: 40px; max-width: 960px; margin: 0 auto;
    line-height: 1.8;
  }}
  h1 {{ font-size: 22px; color: #f97316; border-bottom: 2px solid #f97316; padding-bottom: 8px; margin: 24px 0 12px; }}
  h2 {{ font-size: 18px; color: #fb923c; margin: 20px 0 10px; }}
  h3 {{ font-size: 15px; color: #fbbf24; margin: 16px 0 8px; }}
  p {{ margin: 6px 0; font-size: 14px; }}
  strong {{ color: #fbbf24; }}
  table {{
    width: 100%; border-collapse: collapse; margin: 16px 0;
    font-size: 13px;
  }}
  th, td {{
    border: 1px solid #444; padding: 8px 12px; text-align: left;
  }}
  th {{ background: #2a2a4a; color: #f97316; font-weight: 600; }}
  tr:nth-child(even) td {{ background: #1e1e3a; }}
  .header {{
    background: linear-gradient(135deg, #1e1e3a 0%, #2a1a3e 100%);
    padding: 20px 24px; border-radius: 12px;
    margin-bottom: 24px; border: 1px solid #333;
  }}
  .header h1 {{ border: none; margin: 0; padding: 0; }}
  .header .meta {{ font-size: 12px; color: #888; margin-top: 4px; }}
</style>
</head>
<body>
<div class="header">
  <h1>📄 {filename}</h1>
  <div class="meta">在线预览 · 来源: uploads/client_materials/{path}</div>
</div>
{body}
</body>
</html>"""
            return HTMLResponse(page)

        except Exception as e:
            logger.error(f"DOCX preview failed: {e}")
            return HTMLResponse(f"<h1>DOCX 预览失败</h1><p>{str(e)}</p>", status_code=500)

    # ── PDF → serve inline ──
    elif ext == ".pdf":
        return FileResponse(
            full_path,
            media_type="application/pdf",
            headers={"Content-Disposition": f"inline; filename*=UTF-8''{filename}"},
        )

    # ── Images → serve inline ──
    elif ext in (".png", ".jpg", ".jpeg", ".gif", ".webp"):
        media_types = {
            ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
            ".gif": "image/gif", ".webp": "image/webp",
        }
        return FileResponse(full_path, media_type=media_types.get(ext, "image/png"))

    else:
        # Unsupported: offer download
        return FileResponse(
            full_path,
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
        )


def _sse(data: dict) -> str:
    """Format data as SSE event."""
    return f"data: {json.dumps(data, ensure_ascii=False)}\n\n"

